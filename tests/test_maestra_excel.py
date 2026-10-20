"""The table's «Exportar a Excel» and «Importar Excel».

Through the real HTTP dispatcher with real sessions, on SQLite and (when
ARA_MAP_TEST_DATABASE_URL is set) on a disposable Postgres schema, with the
fictional fixture of tests/test_roles_y_bases.py: administrators ada and alan;
operators olga (Base Uno), omar (Base Uno and Base Dos), otto (no grant).
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from unittest.mock import patch

from server import postgres
from tests.test_roles_y_bases import URL, Escenario, filas

CSV = (
    "Terreno,Estado,Municipio,Superficie m2,Asking Price,Tipo de terreno,Comentarios,Corredor\n"
    "Lote Excel Uno,Jalisco,Zapopan,1200,350000,Industrial,Nota ficticia,Ana\n"
    "Lote Excel Dos,Sonora,Hermosillo,800,,Agrícola,,\n"
    ",Jalisco,Sin nombre,10,,,,\n"
).encode()


class Excel(Escenario):
    """The checks. Mixed into one TestCase per backend."""

    def crudo(self, method, path, data=None, user="ada", headers=None):
        peticion = urllib.request.Request(
            f"http://127.0.0.1:{self.httpd.server_address[1]}{path}", method=method, data=data,
            headers={"Cookie": self.cookies[user], **(headers or {})})
        try:
            with urllib.request.urlopen(peticion, timeout=20) as respuesta:
                return respuesta.status, respuesta.headers, respuesta.read()
        except urllib.error.HTTPError as error:
            with error:
                return error.code, error.headers, error.read()

    def importar(self, base, contenido=CSV, simular=False, user="ada", nombre="terrenos.csv"):
        estado, _, cuerpo = self.crudo(
            "POST", f"/api/maestra/bases/{base}/importar" + ("?simular=1" if simular else ""),
            contenido, user, {"Content-Type": "application/octet-stream",
                              "X-Archivo": urllib.parse.quote(nombre)})
        return estado, json.loads(cuerpo)

    def exportar(self, path, user="ada"):
        estado, cabeceras, cuerpo = self.crudo("GET", path, user=user)
        if estado != 200:
            return estado, json.loads(cuerpo)
        from openpyxl import load_workbook  # vendored: reachable once server is imported
        self.assertIn("spreadsheetml", cabeceras["Content-Type"])
        hoja = load_workbook(io.BytesIO(cuerpo)).active
        return estado, [list(r) for r in hoja.iter_rows(values_only=True)]

    def de_base(self, base):
        return {r["terreno"]: r for r in filas(
            "inventory_revision", "id IN (SELECT draft_revision_id FROM inventory_terrain"
            " WHERE base_id = ?)", (base,))}

    # -- import --------------------------------------------------------------

    def test_a_dry_run_reports_and_writes_nothing(self):
        antes = self.cuenta("inventory_terrain")
        estado, r = self.importar(self.b1, simular=True)
        self.assertEqual(estado, 200, r)
        self.assertEqual((r["simulacion"], r["leidas"], r["nuevas"], r["total_duplicadas"]),
                         (True, 2, 2, 0))
        self.assertEqual(len(r["sin_nombre"]), 1)
        self.assertEqual(r["no_cargadas"], ["Corredor"])
        self.assertEqual(self.cuenta("inventory_terrain"), antes)

    def test_an_import_creates_the_rows_with_type_and_comments_and_a_rerun_skips_them(self):
        estado, r = self.importar(self.b1)
        self.assertEqual((estado, r["nuevas"]), (200, 2), r)
        filas_b1 = self.de_base(self.b1)
        uno = filas_b1["Lote Excel Uno"]
        self.assertEqual((uno["tipo_terreno"], uno["notas_internas"], uno["asking_price"],
                          uno["superficie_m2"]), ("Industrial", "Nota ficticia", 350000, 1200))
        self.assertIsNone(filas_b1["Lote Excel Dos"]["asking_price"])     # no price is not zero
        self.assertEqual(self.cuenta("inventory_event", "action = 'create'"), 5 + 2)
        estado, r = self.importar(self.b1)
        self.assertEqual((estado, r["nuevas"], r["total_duplicadas"]), (200, 0, 2), r)
        self.assertEqual(len(self.de_base(self.b1)), 2 + 2)              # t1, ta and the two

    def test_one_invalid_value_stops_the_whole_import(self):
        malo = CSV + b"Lote Excel Tres,Jalisco,Zapopan,-5,,,,\n"
        antes = self.cuenta("inventory_terrain")
        estado, r = self.importar(self.b1, malo)
        self.assertEqual(estado, 422, r)
        self.assertEqual(r["detalle"]["total_errores"], 1)
        self.assertIn("Lote Excel Tres", r["detalle"]["errores"][0])
        self.assertEqual(self.cuenta("inventory_terrain"), antes)

    def test_only_administrators_import_and_only_into_an_existing_base(self):
        antes = self.cuenta("inventory_terrain")
        self.assertEqual(self.importar(self.b1, user="olga")[0], 403)
        self.assertEqual(self.importar(str(uuid.uuid4()))[0], 404)
        self.assertEqual(self.importar(self.b1, b"x", nombre="notas.txt")[0], 400)
        self.assertEqual(self.cuenta("inventory_terrain"), antes)

    # -- export --------------------------------------------------------------

    def test_a_base_exports_its_own_rows_with_the_filters_and_reimports_cleanly(self):
        self.importar(self.b1)
        estado, hoja = self.exportar(f"/api/maestra/bases/{self.b1}/exportar?estado=Jalisco",
                                     user="olga")
        self.assertEqual(estado, 200, hoja)
        self.assertEqual(hoja[0][:5], ["Tipo de terreno", "Terreno", "Estado", "Municipio",
                                       "Superficie m2"])
        self.assertNotIn("Base", hoja[0])                  # one base: no Base column
        self.assertEqual([r[1] for r in hoja[1:]], ["Lote Excel Uno"])
        self.assertEqual(hoja[1][hoja[0].index("Comentarios")], "Nota ficticia")
        # What a base exports, another base imports as the same rows.
        estado, cabeceras, libro = self.crudo("GET", f"/api/maestra/bases/{self.b1}/exportar")
        estado, r = self.importar(self.b2, libro, nombre="terrenos.xlsx")
        self.assertEqual((estado, r["nuevas"], r["no_cargadas"]), (200, 4, []), r)
        uno = self.de_base(self.b2)["Lote Excel Uno"]
        self.assertEqual((uno["tipo_terreno"], uno["notas_internas"]), ("Industrial", "Nota ficticia"))

    def test_the_master_export_is_for_administrators_and_names_each_base(self):
        self.importar(self.b1)
        estado, hoja = self.exportar("/api/maestra/exportar?q=Excel")
        self.assertEqual(estado, 200, hoja)
        base = hoja[0].index("Base")
        self.assertEqual(sorted((r[1], r[base]) for r in hoja[1:]),
                         [("Lote Excel Dos", "Base Uno"), ("Lote Excel Uno", "Base Uno")])
        self.assertEqual(self.exportar("/api/maestra/exportar", user="olga")[0], 403)
        self.assertEqual(self.exportar(f"/api/maestra/bases/{self.b1}/exportar", user="otto")[0], 404)
        self.assertEqual(self.exportar("/api/maestra/exportar?sort=nada")[0], 422)


class ExcelSqlite(Excel, unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self._entorno = patch.dict(os.environ, {"ARA_MAP_DB": str(Path(self._dir.name) / "prueba.db")})
        self._entorno.start()
        os.environ.pop("ARA_MAP_DATABASE_URL", None)
        self.addCleanup(self._dir.cleanup)
        self.addCleanup(self._entorno.stop)
        self.levantar()


@unittest.skipUnless(URL, "No Postgres test connection configured")
class ExcelPostgres(Excel, unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import psycopg
        from psycopg.conninfo import make_conninfo
        cls.schema = "test_ara_excel_" + uuid.uuid4().hex
        with psycopg.connect(URL) as conn:
            conn.execute(f'CREATE SCHEMA "{cls.schema}"')
        cls.env = patch.dict(os.environ, {"ARA_MAP_DATABASE_URL": make_conninfo(
            URL, options=f"-c search_path={cls.schema}")})
        cls.env.start()
        with postgres.session() as conn:
            conn.raw.execute(postgres.schema_sql(), prepare=False)
            postgres.migrate(conn)

    @classmethod
    def tearDownClass(cls):
        import psycopg
        cls.env.stop()
        with psycopg.connect(URL) as conn:
            conn.execute(f'DROP SCHEMA "{cls.schema}" CASCADE')

    def setUp(self):
        with postgres.session() as conn:
            conn.raw.execute("TRUNCATE team_user, team_login_failure CASCADE")
        self.levantar()


if __name__ == "__main__":
    unittest.main()
