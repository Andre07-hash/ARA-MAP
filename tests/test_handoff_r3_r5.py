"""Regressions for the handoff requests R3 (per-source sheets) and R5 (names).

R1 (typing focus) and R4 (double activation) are browser behaviours and are
covered in tests/e2e/smoke.mjs; the geometry they rely on is in tests/js.
"""

from __future__ import annotations

import io
import json
import unittest

from openpyxl import load_workbook

from server import db
from server.api import exportar as api_exportar
from server.api import mapas as api_mapas
from server.importer import TerrainRecord
from server.repo import bases, mapas, terrenos
from server.router import Request
from server.web_util import ApiError
from tests.support import TempDatabase


def req(**body) -> Request:
    params = body.pop("params", {})
    return Request("POST", "/x", {}, params, json.dumps(body).encode(), {})


def rec(orden=1, nombre="El Mirador", **kw) -> TerrainRecord:
    base = dict(
        orden=orden, fila=orden + 1, terreno=nombre, estado="Jalisco",
        municipio="Zapopan", superficie_m2=10_000.0, superficie_ha=1.0,
        asking_price=5_000_000.0, asking_m2=500.0, lat=20.7, lon=-103.4,
    )
    return TerrainRecord(**{**base, **kw})


def hojas(payload: bytes) -> dict[str, int]:
    libro = load_workbook(io.BytesIO(payload))
    return {ws.title: ws.max_row - 1 for ws in libro.worksheets}


class SheetTitles(unittest.TestCase):
    """R3 - layer names are user text and must be made legal without changing."""

    def titulo(self, nombre, usados=None, indice=0):
        return api_exportar.sheet_title(nombre, usados if usados is not None else set(), indice)

    def test_a_plain_name_is_kept(self):
        self.assertEqual(self.titulo("Base Norte"), "Base Norte")

    def test_forbidden_characters_are_replaced(self):
        for prohibido in ":\\/?*[]":
            with self.subTest(prohibido=prohibido):
                titulo = self.titulo(f"Norte{prohibido}Sur")
                self.assertNotIn(prohibido, titulo)

    def test_control_characters_are_removed(self):
        self.assertEqual(self.titulo("Nor\x00te\x1f"), "Norte")

    def test_a_long_name_is_truncated_to_the_excel_limit(self):
        titulo = self.titulo("B" * 60)
        self.assertEqual(len(titulo), 31)

    def test_collisions_get_deterministic_suffixes(self):
        usados = set()
        self.assertEqual(self.titulo("Norte", usados), "Norte")
        self.assertEqual(self.titulo("Norte", usados), "Norte (2)")
        self.assertEqual(self.titulo("Norte", usados), "Norte (3)")

    def test_collisions_are_case_insensitive(self):
        usados = set()
        self.assertEqual(self.titulo("Norte", usados), "Norte")
        self.assertEqual(self.titulo("norte", usados), "norte (2)")

    def test_a_suffix_never_pushes_past_the_limit(self):
        usados = set()
        largo = "B" * 40
        primero = self.titulo(largo, usados)
        segundo = self.titulo(largo, usados)
        self.assertEqual(len(primero), 31)
        self.assertLessEqual(len(segundo), 31)
        self.assertTrue(segundo.endswith(" (2)"))
        self.assertNotEqual(primero, segundo)

    def test_an_empty_name_falls_back_deterministically(self):
        self.assertEqual(self.titulo("   ", indice=2), "Capa 3")
        self.assertEqual(self.titulo(None, indice=0), "Capa 1")

    def test_a_name_that_sanitises_to_nothing_falls_back(self):
        self.assertEqual(self.titulo("///", indice=1), "Capa 2")

    def test_the_reserved_history_name_is_prefixed(self):
        self.assertNotEqual(self.titulo("History").casefold(), "history")

    def test_accents_are_preserved(self):
        self.assertEqual(self.titulo("Comparación año"), "Comparación año")


class ComparisonExport(TempDatabase):
    """R3 - one worksheet per saved layer."""

    def setUp(self):
        super().setUp()
        self.a = bases.create(self.conn, "Prueba 08", None, None)
        self.b = bases.create(self.conn, "Prueba 09", None, None)
        terrenos.insert(self.conn, self.a, [rec(i, f"A{i}") for i in range(1, 13)], {})
        terrenos.insert(self.conn, self.b, [rec(i, f"B{i}") for i in range(1, 14)], {})
        self.ma = mapas.create(self.conn, "Prueba 08", "simple",
                               [{"base_id": self.a, "color": "#1"}], sigue_base=True)
        self.mb = mapas.create(self.conn, "Prueba 09", "simple",
                               [{"base_id": self.b, "color": "#2"}], sigue_base=True)
        self.comp = mapas.merge(self.conn, "Comparación", [self.ma, self.mb], ["#1", "#2"])

    def export(self, **body):
        return api_exportar.export(req(**body))[0]

    def test_each_source_gets_its_own_sheet(self):
        self.assertEqual(
            hojas(self.export(mapa_id=self.comp, nombre="c")),
            {"Prueba 08": 12, "Prueba 09": 13},
        )

    def test_sheet_order_follows_layer_order(self):
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.comp, nombre="c")))
        self.assertEqual([ws.title for ws in libro.worksheets], ["Prueba 08", "Prueba 09"])

    def test_rows_land_in_the_sheet_of_their_own_layer(self):
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.comp, nombre="c")))
        for titulo, prefijo in (("Prueba 08", "A"), ("Prueba 09", "B")):
            nombres = [c[1].value for c in libro[titulo].iter_rows(min_row=2)]
            self.assertTrue(all(n.startswith(prefijo) for n in nombres), titulo)

    def test_a_layer_excluded_by_filters_keeps_a_header_only_sheet(self):
        filas = mapas.terrenos(self.conn, self.comp)
        solo_primera = [f["id"] for f in filas if f["capa"] == 0]
        self.assertEqual(
            hojas(self.export(mapa_id=self.comp, ids=solo_primera, nombre="c")),
            {"Prueba 08": 12, "Prueba 09": 0},
        )

    def test_excluded_rows_do_not_leak_into_another_sheet(self):
        filas = mapas.terrenos(self.conn, self.comp)
        solo_primera = [f["id"] for f in filas if f["capa"] == 0]
        libro = load_workbook(
            io.BytesIO(self.export(mapa_id=self.comp, ids=solo_primera, nombre="c")))
        self.assertEqual(libro["Prueba 09"].max_row, 1)

    def test_a_simple_saved_map_keeps_one_sheet(self):
        self.assertEqual(hojas(self.export(mapa_id=self.ma, nombre="s")),
                         {"Registro Análisis": 12})

    def test_a_live_base_keeps_one_sheet(self):
        self.assertEqual(hojas(self.export(base_id=self.a, nombre="b")),
                         {"Registro Análisis": 12})

    def test_two_versions_of_one_source_get_two_sheets(self):
        terrenos.insert(self.conn, self.a, [rec(99, "Nuevo")], {})
        v2 = mapas.create(self.conn, "Prueba 08 v2", "simple",
                          [{"base_id": self.a, "color": "#1"}])
        historia = mapas.merge(self.conn, "Historia", [self.ma, v2], ["#1", "#2"])
        resultado = hojas(self.export(mapa_id=historia, nombre="h"))
        self.assertEqual(len(resultado), 2)
        self.assertEqual(sorted(resultado.values()), [12, 13])

    def test_an_empty_selection_is_an_error_not_an_empty_download(self):
        with self.assertRaises(ApiError):
            self.export(mapa_id=self.comp, ids=[], nombre="c")

    def test_the_header_row_matches_the_importer(self):
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.comp, nombre="c")))
        encabezados = [c.value for c in libro.worksheets[0][1]]
        self.assertEqual(encabezados, [e for e, _, _ in api_exportar.COLUMNS])
        self.assertEqual(encabezados[11], "X")
        self.assertEqual(encabezados[12], "Y")

    def test_panes_are_frozen_below_the_header(self):
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.comp, nombre="c")))
        for ws in libro.worksheets:
            self.assertEqual(ws.freeze_panes, "A2")

    def test_there_is_no_leftover_default_sheet(self):
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.comp, nombre="c")))
        self.assertNotIn("Sheet", libro.sheetnames)

    def test_numbers_stay_numbers_and_blanks_stay_blank(self):
        terrenos.insert(self.conn, self.b, [rec(50, "Sin coords", lat=None, lon=None)], {})
        mapas.refresh_snapshot(self.conn, self.mb)
        comp = mapas.merge(self.conn, "C2", [self.ma, self.mb], ["#1", "#2"])
        libro = load_workbook(io.BytesIO(self.export(mapa_id=comp, nombre="c")))
        ws = libro["Prueba 09"]
        fila = next(r for r in ws.iter_rows(min_row=2) if r[1].value == "Sin coords")
        self.assertIsInstance(fila[5].value, (int, float))   # Superficie m2
        self.assertIsNone(fila[11].value)                    # X
        self.assertIsNone(fila[12].value)                    # Y

    def test_a_name_beginning_with_equals_is_written_as_text(self):
        terrenos.insert(self.conn, self.a, [rec(60, "=SUMA(A1:A2)")], {})
        mapas.refresh_snapshot(self.conn, self.ma)
        libro = load_workbook(io.BytesIO(self.export(mapa_id=self.ma, nombre="s")))
        celda = next(r[1] for r in libro.worksheets[0].iter_rows(min_row=2)
                     if str(r[1].value).startswith("="))
        self.assertEqual(celda.data_type, "s")

    def test_a_layer_whose_source_was_deleted_still_exports(self):
        bases.delete(self.conn, self.a)
        self.assertEqual(
            hojas(self.export(mapa_id=self.comp, nombre="c")),
            {"Prueba 08": 12, "Prueba 09": 13},
        )


class NamePropagation(TempDatabase):
    """R5 - a renamed source must appear renamed everywhere that shows it."""

    def setUp(self):
        super().setUp()
        self.base_id = bases.create(self.conn, "Prueba 08", None, None)
        terrenos.insert(self.conn, self.base_id, [rec(1, "Uno")], {})
        self.sigue = mapas.create(
            self.conn, "Prueba 08", "simple",
            [{"base_id": self.base_id, "color": "#1"}], sigue_base=True)
        self.propio = mapas.create(
            self.conn, "Opciones cliente López", "simple",
            [{"base_id": self.base_id, "color": "#1"}])

    def nombre_de(self, mapa_id):
        return mapas.get(self.conn, mapa_id)["nombre"]

    def capas_de(self, mapa_id):
        return [c["nombre"] for c in mapas.get(self.conn, mapa_id)["capas"]]

    def test_a_following_title_changes_with_its_source(self):
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.assertEqual(self.nombre_de(self.sigue), "Base Renombrada")

    def test_a_custom_title_is_left_alone(self):
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.assertEqual(self.nombre_de(self.propio), "Opciones cliente López")

    def test_layer_labels_update_on_both_maps(self):
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.assertEqual(self.capas_de(self.sigue), ["Base Renombrada"])
        self.assertEqual(self.capas_de(self.propio), ["Base Renombrada"])

    def test_renaming_a_map_turns_following_off(self):
        mapas.update(self.conn, self.sigue, nombre="Título propio")
        bases.rename(self.conn, self.base_id, "Otra vez")
        self.assertEqual(self.nombre_de(self.sigue), "Título propio")

    def test_version_labels_survive_a_source_rename(self):
        terrenos.insert(self.conn, self.base_id, [rec(2, "Dos")], {})
        v2 = mapas.create(self.conn, "Prueba 08 v2", "simple",
                          [{"base_id": self.base_id, "color": "#1"}])
        historia = mapas.merge(self.conn, "Historia", [self.sigue, v2], ["#1", "#2"])

        bases.rename(self.conn, self.base_id, "Base Renombrada")
        capas = self.capas_de(historia)
        self.assertEqual(capas, ["Base Renombrada · Prueba 08", "Base Renombrada · Prueba 08 v2"])
        self.assertEqual(len(set(capas)), 2, "the two versions became indistinguishable")

    def test_a_comparison_title_is_never_rewritten(self):
        otra = bases.create(self.conn, "Prueba 09", None, None)
        terrenos.insert(self.conn, otra, [rec(1, "Otro")], {})
        mb = mapas.create(self.conn, "Prueba 09", "simple", [{"base_id": otra, "color": "#2"}])
        comp = mapas.merge(self.conn, "Prueba 08 vs Prueba 09", [self.sigue, mb], ["#1", "#2"])

        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.assertEqual(self.nombre_de(comp), "Prueba 08 vs Prueba 09")
        self.assertIn("Base Renombrada", self.capas_de(comp)[0])

    def test_refresh_recovers_a_stale_label(self):
        # Simulate a map saved before renames propagated.
        self.conn.execute(
            "UPDATE mapa_capa SET base_nombre = 'Nombre viejo' WHERE mapa_id = ?",
            (self.sigue,))
        mapas.refresh_snapshot(self.conn, self.sigue)
        self.assertEqual(self.capas_de(self.sigue), ["Prueba 08"])

    def test_renaming_does_not_touch_the_frozen_rows(self):
        antes = [(t["terreno"], t["asking_price"]) for t in mapas.terrenos(self.conn, self.sigue)]
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        despues = [(t["terreno"], t["asking_price"]) for t in mapas.terrenos(self.conn, self.sigue)]
        self.assertEqual(antes, despues)

    def test_renaming_does_not_move_the_refresh_timestamp(self):
        antes = mapas.get(self.conn, self.sigue)["actualizado_en"]
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.assertEqual(mapas.get(self.conn, self.sigue)["actualizado_en"], antes)

    def test_an_orphaned_layer_is_never_relinked_by_name(self):
        bases.delete(self.conn, self.base_id)
        repuesto = bases.create(self.conn, "Prueba 08", None, None)
        terrenos.insert(self.conn, repuesto, [rec(9, f"Ajeno {i}") for i in range(1, 6)], {})

        bases.rename(self.conn, repuesto, "Prueba 08")
        capa = mapas.get(self.conn, self.sigue)["capas"][0]
        self.assertIsNone(capa["base_id"])
        self.assertEqual(capa["conteo"], 1, "the orphan absorbed an unrelated source")

    def test_a_failed_rename_leaves_everything_as_it_was(self):
        original = self.nombre_de(self.sigue)
        try:
            with db.transaction(self.conn):
                bases.rename(self.conn, self.base_id, "A medias")
                raise RuntimeError("falla simulada")
        except RuntimeError:
            pass
        self.assertEqual(bases.get(self.conn, self.base_id)["nombre"], "Prueba 08")
        self.assertEqual(self.nombre_de(self.sigue), original)

    def test_the_names_persist_across_a_reconnect(self):
        bases.rename(self.conn, self.base_id, "Base Renombrada")
        self.conn.close()
        self.conn = db.connect(self.db_path)
        self.assertEqual(self.nombre_de(self.sigue), "Base Renombrada")
        self.assertEqual(self.capas_de(self.sigue), ["Base Renombrada"])

    def test_the_api_reports_what_it_changed(self):
        from server.api import bases as api_bases
        salida = api_bases.rename(
            req(params={"id": self.base_id}, nombre="Base Renombrada"))
        self.assertEqual(salida["base"]["nombre"], "Base Renombrada")
        self.assertEqual(salida["capas_actualizadas"], 2)
        self.assertTrue(any(m["nombre"] == "Base Renombrada" for m in salida["mapas"]))

    def test_the_title_mode_round_trips_through_the_api(self):
        salida = api_mapas.create(req(
            nombre="Prueba 08", tipo="simple",
            capas=[{"base_id": self.base_id, "color": "#1"}],
            nombre_sigue_base=True))
        self.assertTrue(salida["mapa"]["nombre_sigue_base"])


if __name__ == "__main__":
    unittest.main()
