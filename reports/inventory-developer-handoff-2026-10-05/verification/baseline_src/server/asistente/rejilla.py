"""Parse a file's structure into a lossless grid, knowing nothing about terrains.

The legacy readers refuse a table whose header lacks «Terreno». Here structure
and meaning are separate: any well-formed table reaches the mapping stage, and
only structural faults (bad quoting, unsupported encoding, an unreadable or
oversized workbook) are errors.

A CSV can be read with a comma, a semicolon or a tab. When more than one gives
a consistent table, each reading becomes one candidate table (``csv:,`` /
``csv:;`` / ``csv:<tab>``) and the choice is made -- or asked -- like choosing
between worksheets.

Cells keep their original types: numbers typed in Excel stay numbers, text stays
text. Blank rows are dropped, but every kept row carries its source row or
physical line number, so nothing is renumbered.

One piece of Excel presentation is not presentation at all: a cell formatted as
``"USD "#,##0.00`` states the currency of its amount, and that number by itself
does not. Those markers are read per cell and travel with the grid -- into the
packed draft as well -- so the mapping stage can refuse to relabel dollars as
pesos. They are evidence about an amount, never permission to convert it.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import zipfile
import zlib
from collections import Counter
from collections.abc import Sequence
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..csv_importer import DELIMITERS, CsvError, _directive, _read, decode
from ..errors import WorkbookError
from .campos import moneda_formato

PARSER_VERSION = 2

MAX_FILAS = 50_000
MAX_COLUMNAS = 200
MAX_HOJAS = 30
MAX_CELDA = 5_000
MAX_MONEDAS = 5_000   # currency-formatted cells kept per sheet; the rest as column facts
MAX_DESCOMPRIMIDO = 300 * 1024 * 1024  # a workbook is a zip: bound what it expands to

Celda = Any  # str | int | float | bool | None


class EstructuraError(WorkbookError):
    """The file cannot be read as a table at all. Message is for the user."""


@dataclass(frozen=True)
class Hoja:
    """One candidate table: a worksheet, or one reading of a CSV."""

    id: str                              # "sheet:0", "csv:,", "csv:;"
    nombre: str                          # worksheet name, or "CSV"
    filas: tuple[tuple[Celda, ...], ...]  # nonblank rows, cells as typed
    lineas: tuple[int, ...]              # source row / first physical line of each
    separador: str | None = None
    formulas_sin_valor: tuple[str, ...] = ()  # cells like "C5" with no cached result
    # (row, column, ISO code) for each cell whose Excel number format declares a
    # foreign currency, and the columns whose evidence exceeded MAX_MONEDAS.
    monedas: tuple[tuple[int, int, str], ...] = ()
    monedas_truncadas: tuple[tuple[int, str], ...] = ()

    @property
    def ancho(self) -> int:
        return max((len(f) for f in self.filas), default=0)

    def monedas_de(self, posicion: int, filas: Sequence[int]) -> tuple[str, ...]:
        """Foreign currencies the cell formats declare in one column's rows.

        A column whose evidence was truncated counts as carrying it throughout:
        a sheet too large to record cell by cell must not look free of it.
        """
        dentro = set(filas)
        codigos = [c for f, p, c in self.monedas if p == posicion and f in dentro]
        codigos += [c for p, c in self.monedas_truncadas if p == posicion]
        return tuple(dict.fromkeys(codigos))

    def letra(self, posicion: int) -> str:
        """How the user finds a column: Excel letter, or its CSV position."""
        if self.separador is not None:
            return f"columna {posicion + 1}"
        return f"columna {get_column_letter(posicion + 1)}"


@dataclass(frozen=True)
class Rejilla:
    formato: str        # "csv", "xlsx", "xlsm"
    archivo: str
    sha256: str
    hojas: tuple[Hoja, ...]
    version: int = PARSER_VERSION

    def hoja(self, hoja_id: str) -> Hoja:
        for hoja in self.hojas:
            if hoja.id == hoja_id:
                return hoja
        raise KeyError(hoja_id)

    # -- durable storage: a draft must survive separate cloud workers ---------

    def empaquetar(self) -> str:
        data = json.dumps(asdict(self), ensure_ascii=False, separators=(",", ":"))
        return base64.b64encode(zlib.compress(data.encode("utf-8"), 6)).decode("ascii")

    @classmethod
    def desempaquetar(cls, texto: str) -> Rejilla:
        data = json.loads(zlib.decompress(base64.b64decode(texto)).decode("utf-8"))
        hojas = tuple(
            Hoja(
                id=h["id"], nombre=h["nombre"],
                filas=tuple(tuple(f) for f in h["filas"]),
                lineas=tuple(h["lineas"]), separador=h["separador"],
                formulas_sin_valor=tuple(h["formulas_sin_valor"]),
                monedas=tuple((int(f), int(p), str(c)) for f, p, c in h.get("monedas", ())),
                monedas_truncadas=tuple((int(p), str(c)) for p, c in h.get("monedas_truncadas", ())),
            )
            for h in data["hojas"]
        )
        return cls(formato=data["formato"], archivo=data["archivo"], sha256=data["sha256"],
                   hojas=hojas, version=data["version"])


def leer(contenido: bytes, archivo: str) -> Rejilla:
    """Parse an upload by its extension. Raises EstructuraError for the user."""
    sufijo = archivo.lower().rsplit(".", 1)[-1] if "." in archivo else ""
    huella = hashlib.sha256(contenido).hexdigest()
    if sufijo == "csv":
        return Rejilla("csv", archivo, huella, _leer_csv(contenido))
    if sufijo in ("xlsx", "xlsm"):
        return Rejilla(sufijo, archivo, huella, _leer_libro(contenido))
    raise EstructuraError(
        "Formato no admitido. Selecciona un archivo .xlsx, .xlsm o .csv. "
        "Si usas Numbers, expórtalo a Excel o CSV UTF-8."
    )


# ---------------------------------------------------------------------- CSV

def _leer_csv(contenido: bytes) -> tuple[Hoja, ...]:
    try:
        texto = decode(contenido)
    except CsvError as exc:
        raise EstructuraError(str(exc)) from None
    if not texto.strip():
        raise EstructuraError("El CSV está vacío.")

    forzado, linea_directiva = _directive(texto)
    lecturas = [_read(texto, sep, linea_directiva) for sep in ((forzado,) if forzado else DELIMITERS)]
    validas = [lec for lec in lecturas if lec.error is None and lec.rows]
    if not validas:
        falla = next((lec for lec in lecturas if lec.error), None)
        raise EstructuraError(falla.error if falla and falla.error else "El CSV está vacío.")

    candidatas = _lecturas_plausibles(validas)
    hojas = []
    for lectura in candidatas:
        _limitar(len(lectura.rows), max(len(r.values) for r in lectura.rows), "El CSV")
        hojas.append(Hoja(
            id=f"csv:{lectura.delimiter}", nombre="CSV",
            filas=tuple(tuple(_celda_csv(v) for v in r.values) for r in lectura.rows),
            lineas=tuple(r.linea for r in lectura.rows),
            separador=lectura.delimiter,
        ))
    return tuple(hojas)


def _consistencia(lectura: Any) -> tuple[int, float]:
    """(most common row width, share of rows having it)."""
    anchos = Counter(len(r.values) for r in lectura.rows)
    ancho, veces = anchos.most_common(1)[0]
    return ancho, veces / len(lectura.rows)


def _lecturas_plausibles(validas: list[Any]) -> list[Any]:
    """Keep the readings that look like a table; all of them if genuinely unclear.

    Chosen by shape alone -- never by recognizing a terrain header -- so an
    unfamiliar table is not rejected here.
    """
    if len(validas) == 1:
        return validas
    # A file with no separator at all reads identically every way: one table.
    unicas: list[Any] = []
    for lectura in validas:
        if not any(lectura.rows == otra.rows for otra in unicas):
            unicas.append(lectura)
    if len(unicas) == 1:
        return unicas
    con_columnas = [lec for lec in unicas if _consistencia(lec)[0] > 1]
    if len(con_columnas) == 1:          # exactly one reading splits into columns
        return con_columnas
    candidatas = con_columnas or unicas
    consistentes = [lec for lec in candidatas if _consistencia(lec)[1] == 1.0]
    if len(consistentes) == 1:          # exactly one is the same width throughout
        return consistentes
    # Still unclear: best first, the others offered as alternatives.
    return sorted(candidatas, key=lambda lec: _consistencia(lec)[::-1], reverse=True)


def _celda_csv(valor: str) -> Celda:
    if len(valor) > MAX_CELDA:
        raise EstructuraError(
            f"Una celda del archivo supera {MAX_CELDA} caracteres; revisa que las comillas cierren."
        )
    return valor


# -------------------------------------------------------------------- Excel

def _leer_libro(contenido: bytes) -> tuple[Hoja, ...]:
    _revisar_zip(contenido)
    try:
        valores = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        formulas = load_workbook(io.BytesIO(contenido), data_only=False, read_only=True)
    except Exception as exc:  # openpyxl raises a variety of types
        raise EstructuraError(f"No se pudo abrir el archivo de Excel: {exc}") from exc

    try:
        if len(valores.sheetnames) > MAX_HOJAS:
            raise EstructuraError(f"El libro tiene más de {MAX_HOJAS} hojas.")
        hojas = []
        for indice, nombre in enumerate(valores.sheetnames):
            hoja = _leer_hoja(valores[nombre], formulas[nombre], f"sheet:{indice}", str(nombre))
            if hoja.filas:
                hojas.append(hoja)
    finally:
        valores.close()
        formulas.close()
    if not hojas:
        raise EstructuraError("El libro no contiene ninguna hoja con datos.")
    return tuple(hojas)


def _revisar_zip(contenido: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
            total = sum(info.file_size for info in archivo.infolist())
    except zipfile.BadZipFile:
        raise EstructuraError("No se pudo abrir el archivo de Excel: no es un libro .xlsx válido.") from None
    if total > MAX_DESCOMPRIMIDO:
        raise EstructuraError("El libro es demasiado grande una vez descomprimido para importarlo.")


def _leer_hoja(valores: Any, formulas: Any, hoja_id: str, nombre: str) -> Hoja:
    filas: list[tuple[Celda, ...]] = []
    lineas: list[int] = []
    sin_valor: list[str] = []
    monedas: list[tuple[int, int, str]] = []
    truncadas: list[tuple[int, str]] = []
    # Formulas never run: the cached result is read, and a formula whose result
    # was never saved is reported instead of being taken as an empty cell.
    # The value pass reads cells rather than bare values so that each one's
    # number format -- where the currency of an amount is written -- survives.
    pares = zip(valores.iter_rows(), formulas.iter_rows(values_only=True))
    for numero, (fila, crudas) in enumerate(pares, start=1):
        celdas = [_celda_excel(getattr(c, "value", None)) for c in fila]
        for posicion, (valor, cruda) in enumerate(zip(celdas, crudas)):
            if valor is None and isinstance(cruda, str) and cruda.startswith("="):
                sin_valor.append(f"{get_column_letter(posicion + 1)}{numero}")
        while celdas and celdas[-1] is None:
            celdas.pop()
        if not any(c is not None and (not isinstance(c, str) or c.strip()) for c in celdas):
            continue
        _limitar(len(filas) + 1, len(celdas), f"La hoja «{nombre}»")
        _monedas_de_fila(fila, celdas, len(filas), monedas, truncadas)
        filas.append(tuple(celdas))
        lineas.append(numero)
    return Hoja(id=hoja_id, nombre=nombre, filas=tuple(filas), lineas=tuple(lineas),
                formulas_sin_valor=tuple(sin_valor[:50]), monedas=tuple(monedas),
                monedas_truncadas=tuple(dict.fromkeys(truncadas)))


def _monedas_de_fila(fila: Any, celdas: list[Celda], indice: int,
                     monedas: list[tuple[int, int, str]], truncadas: list[tuple[int, str]]) -> None:
    """Record which cells of one kept row are formatted in a foreign currency.

    Only cells that carry a value count: an empty cell formatted in dollars
    says nothing about an amount. Past MAX_MONEDAS the column and code are
    remembered instead, so a very large sheet cannot look free of evidence.
    """
    for posicion, celda in enumerate(fila):
        if posicion >= len(celdas) or celdas[posicion] is None:
            continue
        codigo = moneda_formato(getattr(celda, "number_format", None))
        if codigo is None:
            continue
        if len(monedas) < MAX_MONEDAS:
            monedas.append((indice, posicion, codigo))
        else:
            truncadas.append((posicion, codigo))


def _celda_excel(valor: Any) -> Celda:
    if isinstance(valor, (datetime, date, time)):
        return valor.isoformat()
    if isinstance(valor, str) and len(valor) > MAX_CELDA:
        raise EstructuraError(f"Una celda del libro supera {MAX_CELDA} caracteres.")
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    return str(valor)


def _limitar(filas: int, columnas: int, donde: str) -> None:
    if filas > MAX_FILAS:
        raise EstructuraError(f"{donde} tiene más de {MAX_FILAS:,} filas; divídelo en archivos más pequeños.")
    if columnas > MAX_COLUMNAS:
        raise EstructuraError(f"{donde} tiene más de {MAX_COLUMNAS} columnas.")
