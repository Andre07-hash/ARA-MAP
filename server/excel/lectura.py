"""Read a connected workbook's bytes into keyed, normalized terrain records.

Provider-independent: it receives the downloaded bytes and the source's
configuration (sheet, stable-ID column, currency) and never touches the
database. It reuses the ordinary importer's header aliases and row builder, so
a connected sheet is interpreted exactly like an uploaded one (including the
X = latitude convention).

Rules that differ from an ordinary upload, because a refresh REPLACES the live
view and a dropped row would look like a deletion:

* Every row that carries data must have a stable ID and a terrain name; a
  number that cannot be read, a duplicate ID or an unsupported currency is an
  error. Any error blocks the whole candidate (nothing partial activates).
* IDs are compared exactly as text. Surrounding whitespace is trimmed; case,
  leading zeros and inner characters are kept. A numeric cell 12 / 12.0 is
  the ID "12". Booleans and dates are refused as IDs. An ID is never inferred
  from the row position or the terrain name.
* Formulas are read through the values Excel last saved with the file
  (openpyxl data_only); a formula with no saved value reads as empty.
* The header is the first row of the selected sheet. Merged or multi-row
  headers are not interpreted; the "Terreno" column and the ID column must be
  in that first row.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import io
import json
import math
import zipfile
from dataclasses import dataclass, replace
from typing import Any

from ..asistente.campos import CLAVES_MONEDA, SOPORTADAS, moneda_de_valor
from ..importer import (
    REQUIRED_FIELDS,
    RejectedRow,
    TerrainRecord,
    _build_record,
    map_headers,
    resolve_sheet,
)
from ..normalize import clean_text, fold

# Bump when normalization changes: it is part of every configuration's
# fingerprint, so a parser change re-evaluates sources instead of trusting an
# old interpretation of unchanged bytes.
VERSION_LECTOR = "1"

MAX_BYTES = 10 * 1024 * 1024          # downloaded workbook
MAX_EXPANDIDO = 80 * 1024 * 1024      # sum of uncompressed ZIP members
MAX_MIEMBROS = 2_000
MAX_FILAS = 20_000
MAX_COLUMNAS = 200
MONEDAS = ("USD", "MXN", "desconocida", "columna")

# The normalized business fields of one row, in a fixed order.
CAMPOS = ("terreno", "estado", "municipio", "direccion", "superficie_m2", "superficie_ha",
          "afectaciones_pct", "afectaciones_m2", "asking_price", "asking_m2", "lat", "lon",
          "moneda", "id_origen", "extra")


class LecturaError(Exception):
    """The file cannot be read at all (not a candidate)."""

    def __init__(self, codigo: str, mensaje: str) -> None:
        super().__init__(mensaje)
        self.codigo = codigo
        self.mensaje = mensaje


@dataclass(frozen=True)
class Problema:
    codigo: str
    mensaje: str
    fila: int | None = None

    def dto(self) -> dict[str, Any]:
        return {"codigo": self.codigo, "mensaje": self.mensaje, "fila": self.fila}


@dataclass(frozen=True)
class Fila:
    clave: str
    fila: int
    record: TerrainRecord
    datos: dict[str, Any]
    huella: str


@dataclass(frozen=True)
class Lectura:
    hojas: tuple[str, ...]
    hoja: str
    encabezados: tuple[str, ...]
    filas: tuple[Fila, ...]
    problemas: tuple[Problema, ...]

    @property
    def valida(self) -> bool:
        return not self.problemas

    @property
    def contenido_huella(self) -> str:
        """Order-insensitive fingerprint of the business content."""
        return huella_contenido((f.clave, f.huella) for f in self.filas)


@dataclass(frozen=True)
class Configuracion:
    hoja: str
    columna_id: str
    moneda: str
    columna_moneda: str | None = None
    version_lector: str = VERSION_LECTOR

    def huella(self, drive_id: str, item_id: str) -> str:
        return _sha(json.dumps({"drive": drive_id, "item": item_id, "hoja": self.hoja,
                                "id": self.columna_id, "moneda": self.moneda,
                                "columna_moneda": self.columna_moneda,
                                "lector": self.version_lector}, sort_keys=True, ensure_ascii=False))


def _sha(texto: str) -> str:
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def huella_contenido(pares: Any) -> str:
    return _sha(json.dumps(sorted(pares), ensure_ascii=False))


def huella_datos(datos: dict[str, Any]) -> str:
    return _sha(json.dumps(datos, sort_keys=True, ensure_ascii=False))


def clave_de(raw: Any) -> str | None:
    """A stable ID cell as exact text, or None when empty. Raises ValueError
    for a type that cannot be an ID."""
    if raw is None:
        return None
    if isinstance(raw, (bool, _dt.date, _dt.time, _dt.datetime)):
        raise ValueError("una fecha o casilla de verificación no puede ser un ID")
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if not math.isfinite(raw):
            raise ValueError("número no válido")
        return str(int(raw)) if raw.is_integer() else repr(raw)
    texto = str(raw).strip()
    return texto or None


def comprobar_zip(contenido: bytes) -> None:
    """Refuse oversized or ZIP-bomb-like workbooks before openpyxl expands them."""
    if len(contenido) > MAX_BYTES:
        raise LecturaError("demasiado_grande",
                           f"El archivo pesa más de {MAX_BYTES // (1024 * 1024)} MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            miembros = z.infolist()
    except zipfile.BadZipFile:
        raise LecturaError("archivo_invalido", "El archivo no es un libro de Excel (.xlsx).") from None
    if len(miembros) > MAX_MIEMBROS or sum(m.file_size for m in miembros) > MAX_EXPANDIDO:
        raise LecturaError("demasiado_grande", "El libro es demasiado grande al descomprimirse.")


def _abrir(contenido: bytes) -> Any:
    comprobar_zip(contenido)
    from openpyxl import load_workbook
    try:
        return load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many types
        raise LecturaError("archivo_invalido", "No se pudo abrir el libro de Excel.") from None


def examinar(contenido: bytes, hoja: str | None = None) -> dict[str, Any]:
    """Setup preview: sheets, headers and, per column, how usable it is as a
    stable ID. Reads no configuration and decides nothing."""
    libro = _abrir(contenido)
    try:
        hojas = tuple(str(n) for n in libro.sheetnames)
        nombre = hoja if hoja is not None else resolve_sheet(libro)
        if nombre not in hojas:
            raise LecturaError("hoja_faltante", f"El libro no tiene la hoja «{nombre}».")
        filas = libro[nombre].iter_rows(values_only=True, max_col=MAX_COLUMNAS)
        encabezado = next(filas, ())
        etiquetas = [clean_text(v) for v in encabezado]
        valores: list[list[Any]] = [[] for _ in etiquetas]
        total = 0
        for valores_fila in filas:
            if total >= MAX_FILAS:
                break
            if not any(clean_text(v) is not None for v in valores_fila):
                continue
            total += 1
            for i in range(len(etiquetas)):
                valores[i].append(valores_fila[i] if i < len(valores_fila) else None)
    finally:
        libro.close()
    mapped, _ = map_headers(encabezado)
    columnas = []
    for i, etiqueta in enumerate(etiquetas):
        if etiqueta is None:
            continue
        claves = []
        for v in valores[i]:
            try:
                claves.append(clave_de(v))
            except ValueError:
                claves.append(None)
        llenas = [c for c in claves if c is not None]
        columnas.append({
            "encabezado": etiqueta, "campo": mapped.get(i), "llenas": len(llenas),
            "unicas": len(set(llenas)), "puede_ser_id": len(llenas) == total and len(set(llenas)) == total,
            "es_moneda": fold(etiqueta) in CLAVES_MONEDA,
        })
    return {"hojas": list(hojas), "hoja": nombre, "filas": total, "columnas": columnas,
            "tiene_terreno": all(f in mapped.values() for f in REQUIRED_FIELDS),
            "id_sugerido": next((c["encabezado"] for c in columnas
                                 if c["campo"] == "id_origen" and c["puede_ser_id"]), None)}


def leer(contenido: bytes, config: Configuracion) -> Lectura:
    """The full candidate under one configuration. Problems are collected, not
    raised: the caller decides that any problem blocks activation."""
    if config.moneda not in MONEDAS:
        raise LecturaError("configuracion", "Moneda de configuración no admitida.")
    libro = _abrir(contenido)
    problemas: list[Problema] = []
    filas: list[Fila] = []
    try:
        hojas = tuple(str(n) for n in libro.sheetnames)
        if config.hoja not in hojas:
            raise LecturaError("hoja_faltante", f"El libro ya no tiene la hoja «{config.hoja}».")
        rows = libro[config.hoja].iter_rows(values_only=True, max_col=MAX_COLUMNAS)
        encabezado = tuple(next(rows, ()))
        etiquetas = [clean_text(v) for v in encabezado]
        mapped, unknown = map_headers(encabezado)
        if not all(f in mapped.values() for f in REQUIRED_FIELDS):
            raise LecturaError("encabezado", "La primera fila de la hoja no tiene la columna «Terreno».")
        indice_id = _indice(etiquetas, config.columna_id)
        if indice_id is None:
            raise LecturaError("columna_id_faltante",
                               f"La hoja ya no tiene la columna de ID «{config.columna_id}».")
        indice_moneda = None
        if config.moneda == "columna":
            indice_moneda = _indice(etiquetas, config.columna_moneda or "")
            if indice_moneda is None:
                raise LecturaError("columna_moneda_faltante",
                                   f"La hoja ya no tiene la columna de moneda «{config.columna_moneda}».")
        # The ID and currency columns are configuration, not extra data.
        for indice in (indice_id, indice_moneda):
            if indice is not None:
                unknown.pop(indice, None)
        vistas: dict[str, int] = {}
        for numero, valores in enumerate(rows, start=2):
            if numero - 1 > MAX_FILAS:
                raise LecturaError("demasiado_grande", f"La hoja tiene más de {MAX_FILAS} filas.")
            valores = tuple(valores)
            if not any(clean_text(v) is not None for v in valores):
                continue  # a blank row is not a terrain
            try:
                clave = clave_de(valores[indice_id] if indice_id < len(valores) else None)
            except ValueError as exc:
                problemas.append(Problema("id_invalido", f"ID no válido: {exc}.", numero))
                continue
            if clave is None:
                problemas.append(Problema("id_faltante", "La fila tiene datos pero no tiene ID.", numero))
                continue
            if clave in vistas:
                problemas.append(Problema(
                    "id_duplicado", f"El ID «{clave}» se repite (también en la fila {vistas[clave]}).", numero))
                continue
            vistas[clave] = numero
            resultado = _build_record(valores, mapped, unknown, len(filas) + 1, numero)
            if isinstance(resultado, RejectedRow) or resultado is None:
                problemas.append(Problema("sin_nombre", f"El ID «{clave}» no tiene nombre de terreno.", numero))
                continue
            for nota in resultado.notes:
                problemas.append(Problema("numero_ilegible",
                                          f"«{nota.valor}» no es un número válido en {nota.campo}.", numero))
            moneda, problema = _moneda(config, valores, indice_moneda, resultado, numero)
            if problema:
                problemas.append(problema)
            record = replace(resultado, moneda=moneda, notes=())
            datos = normalizar(record)
            filas.append(Fila(clave, numero, record, datos, huella_datos(datos)))
    finally:
        libro.close()
    return Lectura(hojas, config.hoja, tuple(e for e in etiquetas if e), tuple(filas), tuple(problemas))


def normalizar(record: TerrainRecord) -> dict[str, Any]:
    datos = {c: getattr(record, c) for c in CAMPOS if c != "extra"}
    datos["extra"] = dict(sorted((record.extra or {}).items()))
    return datos


def _indice(etiquetas: list[str | None], nombre: str) -> int | None:
    """Columns are found by their exact header text (trimmed), then by folded
    text, so a header moved to another position still matches."""
    for i, e in enumerate(etiquetas):
        if e is not None and e == nombre.strip():
            return i
    for i, e in enumerate(etiquetas):
        if e is not None and fold(e) == fold(nombre):
            return i
    return None


def _moneda(config: Configuracion, valores: tuple[Any, ...], indice: int | None,
            record: TerrainRecord, numero: int) -> tuple[str | None, Problema | None]:
    if config.moneda in SOPORTADAS:
        return config.moneda, None
    if config.moneda == "desconocida" or indice is None:
        return None, None  # an honest unknown: prices are never relabelled
    texto = clean_text(valores[indice] if indice < len(valores) else None)
    if texto is None:
        return None, None
    moneda = moneda_de_valor(texto)
    if moneda not in SOPORTADAS:
        return None, Problema("moneda_no_admitida",
                              f"Moneda «{texto}» no admitida; usa USD o MXN.", numero)
    return moneda, None
