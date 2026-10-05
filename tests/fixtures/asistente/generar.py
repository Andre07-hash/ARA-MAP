"""Write the import-assistant fixtures. Every record is fictional.

    python3 tests/fixtures/asistente/generar.py

Files are committed; this script documents exactly what is in them and lets
them be regenerated. Expected outcomes live in the tests (and, for the held-out
evaluation set, in evaluacion/esperado.json).
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

AQUI = Path(__file__).parent

PREDIOS = [
    ("Predio Alfa", "Jalisco", "Tala", 25000, 2.5, 12500000, 500, 20.653, -103.701),
    ("Predio Beta", "Estado de México", "Tecámac", 40000, 4.0, 18000000, 450, 19.713, -98.968),
    ("Predio Gama", "Nuevo León", "Apodaca", 12000, 1.2, 9600000, 800, 25.781, -100.188),
    ("Predio Delta", "Querétaro", "El Marqués", 60000, 6.0, 21000000, 350, 20.627, -100.271),
]


def csv(nombre: str, lineas: list[str], sep: str = ",") -> None:
    (AQUI / nombre).write_text("\n".join(lineas) + "\n", encoding="utf-8")


def libro(nombre: str, hojas: list[tuple[str, list[list[object]]]],
          formatos: dict[str, str] | None = None) -> None:
    """`formatos` sets an Excel number format per cell ("B2": '"USD "#,##0.00')."""
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, filas in hojas:
        hoja = wb.create_sheet(titulo)
        for fila in filas:
            hoja.append(fila)
    for celda, formato in (formatos or {}).items():
        wb[wb.sheetnames[0]][celda].number_format = formato
    wb.save(AQUI / nombre)


def main() -> None:
    # 1. Familiar: ARA's own headers, reordered. Upload -> preview.
    csv("familiar_reordenado.csv", [
        "Terreno,Asking Price,Estado,Municipio,Superficie m2,X,Y,ID",
        *[f"{n},{p},{e},{m},{s},{la},{lo},{i}" for i, (n, e, m, s, _, p, _, la, lo) in enumerate(PREDIOS, 1)],
    ])

    # 2. Unfamiliar but unambiguous: title rows, curated synonyms, hectares,
    #    a contact column, a second sheet of notes, and a totals row.
    datos = [["Reporte de terrenos disponibles", None, None],
             ["Actualizado septiembre 2026 (datos ficticios)"],
             ["Nombre comercial", "Entidad", "Municipio", "Área del predio (ha)", "Valor de venta MXN",
              "Latitud", "Longitud", "Contacto"],
             *[[n, e, m, ha, p, la, lo, "Tel. 55 0000 0000"] for n, e, m, _, ha, p, _, la, lo in PREDIOS],
             ["Total", None, None, 13.7, 61100000, None, None, None]]
    libro("desconocido_titulos.xlsx", [("Terrenos", datos),
                                       ("Notas", [["Notas internas"], ["Revisar precios en octubre."]])])

    # 3. Unfamiliar headers detection cannot settle: needs automatic
    #    assistance, or questions when it is unavailable.
    csv("desconocido_ia.csv", [
        "Desarrollo,Edo.,Mpio.,Sup. aprox (m2),Precio pedido MXN,Lat. dec,Lon. dec,Teléfono contacto",
        *[f"{n},{e},{m},{s},{p},{la},{lo},55 0000 000{i}"
          for i, (n, e, m, s, _, p, _, la, lo) in enumerate(PREDIOS)],
    ])

    # 4. Ambiguous: "Valor" could be a total price or a price per m2.
    csv("ambiguo_valor.csv", [
        "Terreno,Estado,Municipio,Superficie m2,Valor,Latitud,Longitud",
        *[f"{n},{e},{m},{s},{p},{la},{lo}" for n, e, m, s, _, p, _, la, lo in PREDIOS],
    ])

    # 5. X/Y written the GIS way round (X = longitude): values contradict ARA's convention.
    csv("coordenadas_gis.csv", [
        "Terreno,Estado,Municipio,X,Y",
        *[f"{n},{e},{m},{lo},{la}" for n, e, m, _, _, _, _, la, lo in PREDIOS],
    ])

    # 6. Numbers whose convention the values cannot settle (1,234 = 1234 or 1.234?).
    csv("decimal_ambiguo.csv", [
        "Terreno;Estado;Superficie m2;Asking Price",
        "Predio Alfa;Jalisco;1,250;3,500",
        "Predio Beta;Jalisco;2,400;4,100",
    ])

    # 7. Prices in dollars must never become MXN.
    csv("dolares.csv", [
        "Terreno,Estado,Superficie m2,Precio (USD)",
        *[f"{n},{e},{s},{p // 18}" for n, e, _, s, _, p, _, _, _ in PREDIOS],
    ])

    # 8. Two "Precio" columns and a data column with no heading.
    libro("duplicados.xlsx", [("Hoja1", [
        ["Terreno", "Precio", "Precio", None, "Estado"],
        *[[n, p, a, "Lote 1", e] for n, e, _, _, _, p, a, _, _ in PREDIOS],
    ])])

    # 9. Two terrain tables in one workbook.
    libro("dos_tablas.xlsx", [
        ("Agosto", [["Terreno", "Estado", "Asking Price"], *[[n, e, p] for n, e, _, _, _, p, _, _, _ in PREDIOS[:2]]]),
        ("Septiembre", [["Terreno", "Estado", "Asking Price"], *[[n, e, p] for n, e, _, _, _, p, _, _, _ in PREDIOS]]),
    ])

    # 10. Norte/Este in projected metres (UTM): not latitude/longitude.
    csv("utm.csv", [
        "Terreno,Estado,Norte,Este",
        "Predio Alfa,Jalisco,2284512,685301",
        "Predio Beta,Jalisco,2281044,690877",
    ])

    # 11. A formula whose result was never saved, and a repeated header row.
    libro("formula_sin_valor.xlsx", [("Terrenos", [
        ["Terreno", "Estado", "Superficie m2", "Asking Price"],
        ["Predio Alfa", "Jalisco", 25000, "=C2*500"],
        ["Terreno", "Estado", "Superficie m2", "Asking Price"],
        ["Predio Beta", "Jalisco", 40000, 18000000],
    ])])

    # 12. Dollars declared only by Excel's cell format: the values look bare.
    libro("moneda_formato.xlsx", [("Terrenos", [
        ["Terreno", "Estado", "Superficie m2", "Precio", "Latitud", "Longitud"],
        *[[n, e, s, p // 18, la, lo] for n, e, _, s, _, p, _, la, lo in PREDIOS],
    ])], formatos={f"D{i}": '"USD "#,##0.00' for i in range(2, 6)})

    # 13. A footer note under the table, and a terrain whose name starts with
    #     "Total": one is not a terrain, the other is.
    csv("nota_al_pie.csv", [
        "Terreno,Estado,Superficie m2,Asking Price,Latitud,Longitud",
        *[f"{'Total ' if i == 1 else ''}{n},{e},{s},{p},{la},{lo}"
          for i, (n, e, _, s, _, p, _, la, lo) in enumerate(PREDIOS, 1)],
        "NOTA: valores sujetos a revisión,,,,,",
    ])

    # 14. A page of notes above the real header: it sits past the bounded search.
    libro("encabezado_tardio.xlsx", [("Terrenos", [
        *[[f"Nota ficticia {i}"] for i in range(1, 56)],
        ["Terreno", "Estado", "Superficie m2", "Asking Price", "Latitud", "Longitud"],
        *[[n, e, s, p, la, lo] for n, e, _, s, _, p, _, la, lo in PREDIOS],
    ])])

    evaluacion()


def evaluacion() -> None:
    """A held-out set, not used while writing the detection rules."""
    carpeta = AQUI / "evaluacion"
    carpeta.mkdir(exist_ok=True)
    casos = {
        "ev1_inmobiliaria.csv": (
            ["Nombre del predio", "Entidad federativa", "Alcaldía", "Superficie total m2",
             "Precio de lista", "Latitude", "Longitude"],
            ["terreno", "estado", "municipio", "superficie_m2", "asking_price", "lat", "lon"]),
        "ev2_broker.csv": (
            ["Propiedad", "Estado", "Municipio", "Hectáreas", "Precio total", "Precio por m2", "Lat", "Lng"],
            ["terreno", "estado", "municipio", "superficie_ha", "asking_price", "asking_m2", "lat", "lon"]),
        "ev3_minimo.csv": (["Predio"], ["terreno"]),
        "ev4_ingles.csv": (
            ["Name", "State", "City", "Area (m2)", "Asking price", "Latitude", "Longitude"],
            ["terreno", "estado", "extra", "superficie_m2", "asking_price", "lat", "lon"]),
        "ev5_abreviado.csv": (
            ["Desarrollo", "Edo", "Mpio", "Sup m2", "Precio", "Lat", "Lon"],
            ["terreno", "estado", "municipio", "superficie_m2", "asking_price", "lat", "lon"]),
        "ev6_folio.csv": (
            ["Folio", "Nombre", "Domicilio", "Metros cuadrados", "Precio $/m2", "Latitud", "Longitud"],
            ["id_origen", "terreno", "direccion", "superficie_m2", "asking_m2", "lat", "lon"]),
    }
    esperado = {}
    valores = {
        "terreno": lambda p: p[0], "estado": lambda p: p[1], "municipio": lambda p: p[2],
        "extra": lambda p: p[2], "superficie_m2": lambda p: p[3], "superficie_ha": lambda p: p[4],
        "asking_price": lambda p: p[5], "asking_m2": lambda p: p[6], "lat": lambda p: p[7],
        "lon": lambda p: p[8], "direccion": lambda p: f"Calle {p[0][-4:]} 10", "id_origen": lambda p: 1,
    }
    for archivo, (encabezados, campos) in casos.items():
        filas = [",".join(encabezados)]
        for i, predio in enumerate(PREDIOS, 1):
            filas.append(",".join(str(i if c == "id_origen" else valores[c](predio)) for c in campos))
        (carpeta / archivo).write_text("\n".join(filas) + "\n", encoding="utf-8")
        esperado[archivo] = dict(zip(encabezados, campos))
    (carpeta / "esperado.json").write_text(json.dumps(esperado, ensure_ascii=False, indent=2) + "\n",
                                           encoding="utf-8")


if __name__ == "__main__":
    main()
