"""Folders for bases and saved maps. Fictional data, temporary databases only.

Every test runs against a throwaway SQLite file (TempDatabase) with the cloud
database variables cleared, so nothing here can reach a shared workspace.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import load_workbook

from server import db
from server.api import bases as api_bases
from server.api import carpetas as api_carpetas
from server.api import exportar as api_exportar
from server.api import importar as api_importar
from server.api import mapas as api_mapas
from server.repo import carpetas as repo
from server.web_util import ApiError
from tests.support import FIXTURE, TempDatabase
from tests.test_api import req

CSV = (b"Terreno,Estado,Municipio,X,Y\n"
       b"Predio Carpeta,Jalisco,Tala,20.6,-103.7\n")


class FolderCase(TempDatabase):
    def setUp(self):
        self._env = patch.dict(os.environ, {"ARA_MAP_DATABASE_URL": "", "DATABASE_URL": ""})
        self._env.start()
        super().setUp()

    def tearDown(self):
        super().tearDown()
        self._env.stop()

    # -- helpers -----------------------------------------------------------

    def folder(self, tipo, nombre):
        return api_carpetas.create(req(json_body={"tipo": tipo, "nombre": nombre}))["carpeta"]

    def folders(self, tipo):
        return api_carpetas.listing(req(query={"tipo": [tipo]}))

    def move_base(self, base_id, carpeta_id):
        return api_bases.move(req(params={"id": base_id}, json_body={"carpeta_id": carpeta_id}))["base"]

    def move_mapa(self, mapa_id, carpeta_id):
        return api_mapas.move(req(params={"id": mapa_id}, json_body={"carpeta_id": carpeta_id}))["mapa"]

    def base(self, nombre="Base"):
        base_id, _ = self.load_fixture(nombre)
        return base_id

    def mapa(self, base_id, nombre="Mapa", **extra):
        return api_mapas.create(req(json_body={
            "nombre": nombre, "tipo": "simple",
            "capas": [{"base_id": base_id, "color": "#2a78d6"}],
            "config": {"basemap": "claro", "filtros": {"estados": ["Jalisco"]}},
            "nombre_sigue_base": True, **extra,
        }))["mapa"]

    def status_of(self, call):
        with self.assertRaises(ApiError) as caught:
            call()
        return caught.exception.status

    def snapshot(self, mapa_id):
        """Everything a folder operation must leave alone on a saved map."""
        rows = lambda sql: [tuple(r) for r in self.conn.execute(sql, (mapa_id,)).fetchall()]  # noqa: E731
        return {
            "mapa": rows("SELECT id, nombre, tipo, creado_en, actualizado_en, config_json,"
                         " nombre_sigue_base FROM mapa WHERE id = ?"),
            "capas": rows("SELECT * FROM mapa_capa WHERE mapa_id = ? ORDER BY orden"),
            "terrenos": rows("SELECT * FROM mapa_terreno WHERE mapa_id = ? ORDER BY id"),
        }


class FolderCrud(FolderCase):
    def test_create_list_rename_delete_for_both_types(self):
        for tipo in ("bases", "mapas"):
            with self.subTest(tipo=tipo):
                carpeta = self.folder(tipo, "  Cliente   Norte ")
                self.assertEqual((carpeta["nombre"], carpeta["tipo"], carpeta["conteo"]),
                                 ("Cliente Norte", tipo, 0))
                listado = self.folders(tipo)
                self.assertEqual([c["nombre"] for c in listado["carpetas"]], ["Cliente Norte"])

                renombrada = api_carpetas.rename(req(params={"id": carpeta["id"]},
                                                     json_body={"nombre": "Cliente Sur"}))["carpeta"]
                self.assertEqual(renombrada["nombre"], "Cliente Sur")

                out = api_carpetas.remove(req(params={"id": carpeta["id"]}))
                self.assertEqual(out, {"eliminada": "Cliente Sur", "tipo": tipo, "trasladados": 0})
                self.assertEqual(self.folders(tipo)["carpetas"], [])

    def test_types_are_separate(self):
        self.folder("bases", "Jalisco")
        self.folder("mapas", "Jalisco")  # the same name once per dashboard is fine
        self.assertEqual(len(self.folders("bases")["carpetas"]), 1)
        self.assertEqual(len(self.folders("mapas")["carpetas"]), 1)

    def test_empty_folders_listed_alphabetically_ignoring_accents(self):
        for nombre in ("Zacatecas", "Ámbito", "beta"):
            self.folder("bases", nombre)
        self.assertEqual([c["nombre"] for c in self.folders("bases")["carpetas"]],
                         ["Ámbito", "beta", "Zacatecas"])

    def test_normalized_duplicates_are_409(self):
        self.folder("bases", "Cliente Norte")
        for nombre in ("cliente norte", "CLIENTE  NÓRTE", " Cliente Norte "):
            with self.subTest(nombre=nombre):
                self.assertEqual(self.status_of(lambda n=nombre: self.folder("bases", n)), 409)

    def test_renaming_onto_another_folder_is_409_but_onto_itself_is_fine(self):
        a = self.folder("mapas", "Norte")
        self.folder("mapas", "Sur")
        self.assertEqual(self.status_of(lambda: api_carpetas.rename(
            req(params={"id": a["id"]}, json_body={"nombre": "sur"}))), 409)
        out = api_carpetas.rename(req(params={"id": a["id"]}, json_body={"nombre": "NORTE"}))
        self.assertEqual(out["carpeta"]["nombre"], "NORTE")

    def test_name_rules(self):
        for nombre in ("", "   ", "x" * 101, None, 5):
            with self.subTest(nombre=nombre):
                self.assertEqual(self.status_of(lambda n=nombre: self.folder("bases", n)), 400)
        self.assertEqual(self.folder("bases", "x" * 100)["nombre"], "x" * 100)

    def test_reserved_navigation_labels(self):
        for tipo, nombre in (("bases", "Sin carpeta"), ("bases", "todas las bases"),
                             ("mapas", "Todos los mapas"), ("mapas", "sin  CARPETA")):
            with self.subTest(tipo=tipo, nombre=nombre):
                self.assertEqual(self.status_of(lambda t=tipo, n=nombre: self.folder(t, n)), 400)

    def test_invalid_type_and_missing_folder(self):
        self.assertEqual(self.status_of(lambda: self.folder("terrenos", "X")), 400)
        self.assertEqual(self.status_of(lambda: api_carpetas.listing(req())), 400)
        self.assertEqual(self.status_of(lambda: api_carpetas.rename(
            req(params={"id": 99}, json_body={"nombre": "X"}))), 404)
        self.assertEqual(self.status_of(lambda: api_carpetas.remove(req(params={"id": 99}))), 404)

    def test_ids_are_not_recycled(self):
        first = self.folder("bases", "Uno")["id"]
        api_carpetas.remove(req(params={"id": first}))
        self.assertGreater(self.folder("bases", "Dos")["id"], first)

    def test_concurrent_duplicate_is_decided_by_the_constraint(self):
        # Simulates a second editor whose preflight check passed too.
        self.folder("bases", "Carrera")
        with patch.object(repo, "_refuse_duplicate", lambda *a, **k: None), \
                self.assertRaises(repo.CarpetaDuplicadaError):
            repo.create(self.conn, "bases", "carrera", "carrera")
        self.assertEqual(len(self.folders("bases")["carpetas"]), 1)


class Placement(FolderCase):
    def test_move_a_base_in_out_and_between(self):
        base_id = self.base()
        a, b = self.folder("bases", "A"), self.folder("bases", "B")
        self.assertIsNone(api_bases.detail(req(params={"id": base_id}))["base"]["carpeta_id"])

        self.assertEqual(self.move_base(base_id, a["id"])["carpeta_id"], a["id"])
        self.assertEqual(self.move_base(base_id, a["id"])["carpeta_id"], a["id"])  # harmless repeat
        self.assertEqual(self.move_base(base_id, b["id"])["carpeta_id"], b["id"])
        self.assertIsNone(self.move_base(base_id, None)["carpeta_id"])
        self.assertEqual(len(api_bases.listing(req())["bases"]), 1)  # moved, never duplicated

    def test_move_a_map_keeps_its_snapshot_and_timestamp(self):
        mapa = self.mapa(self.base())
        antes = self.snapshot(mapa["id"])
        carpeta = self.folder("mapas", "Informes")
        movido = self.move_mapa(mapa["id"], carpeta["id"])
        self.assertEqual(movido["carpeta_id"], carpeta["id"])
        self.assertEqual(movido["actualizado_en"], mapa["actualizado_en"])
        self.assertEqual(self.snapshot(mapa["id"]), antes)

    def test_wrong_type_and_missing_destinations_change_nothing(self):
        base_id = self.base()
        mapa = self.mapa(base_id)
        base_folder = self.folder("bases", "Bases")
        map_folder = self.folder("mapas", "Mapas")
        self.move_base(base_id, base_folder["id"])

        self.assertEqual(self.status_of(lambda: self.move_base(base_id, map_folder["id"])), 400)
        self.assertEqual(self.status_of(lambda: self.move_mapa(mapa["id"], base_folder["id"])), 400)
        self.assertEqual(self.status_of(lambda: self.move_base(base_id, 999)), 404)
        self.assertEqual(api_bases.detail(req(params={"id": base_id}))["base"]["carpeta_id"],
                         base_folder["id"])
        self.assertIsNone(api_mapas.detail(req(params={"id": mapa["id"]}))["mapa"]["carpeta_id"])

    def test_missing_items_are_404(self):
        carpeta = self.folder("bases", "A")
        self.assertEqual(self.status_of(lambda: self.move_base(999, carpeta["id"])), 404)
        self.assertEqual(self.status_of(lambda: self.move_mapa(999, None)), 404)

    def test_destination_must_be_a_positive_integer_or_null(self):
        base_id = self.base()
        for body in ({}, {"carpeta_id": True}, {"carpeta_id": 1.5}, {"carpeta_id": -1},
                     {"carpeta_id": 0}, {"carpeta_id": "1"}):
            with self.subTest(body=body):
                self.assertEqual(self.status_of(lambda b=body: api_bases.move(
                    req(params={"id": base_id}, json_body=b))), 400)

    def test_folder_deleted_between_check_and_write(self):
        # The foreign key, not the preflight, is the last word.
        base_id = self.base()
        with patch.object(repo, "require_folder", lambda *a, **k: None), \
                self.assertRaises(repo.CarpetaNoExisteError):
            repo.assign(self.conn, "bases", base_id, 12345)
        self.assertIsNone(api_bases.detail(req(params={"id": base_id}))["base"]["carpeta_id"])

    def test_assignments_survive_a_restart(self):
        base_id = self.base()
        carpeta = self.folder("bases", "Persistente")
        self.move_base(base_id, carpeta["id"])
        with db.session(self.db_path) as fresh:
            row = fresh.execute("SELECT carpeta_id FROM base WHERE id = ?", (base_id,)).fetchone()
        self.assertEqual(row["carpeta_id"], carpeta["id"])


class Counts(FolderCase):
    def test_counts_are_items_not_terrains_or_layers(self):
        a, b = self.base("Agosto"), self.base("Septiembre")
        carpeta = self.folder("bases", "Clientes")
        self.move_base(a, carpeta["id"])
        comparacion = api_mapas.create(req(json_body={
            "nombre": "Comparación", "tipo": "comparacion",
            "capas": [{"base_id": a, "color": "#111111"}, {"base_id": b, "color": "#222222"}],
        }))["mapa"]
        map_folder = self.folder("mapas", "Comparaciones")
        self.move_mapa(comparacion["id"], map_folder["id"])
        self.mapa(b, "Suelto")

        bases = self.folders("bases")
        self.assertEqual((bases["total"], bases["sin_carpeta"], bases["carpetas"][0]["conteo"]), (2, 1, 1))
        mapas = self.folders("mapas")
        self.assertEqual((mapas["total"], mapas["sin_carpeta"], mapas["carpetas"][0]["conteo"]), (2, 1, 1))


class FolderDeletion(FolderCase):
    def test_deleting_a_folder_unfiles_its_items_and_keeps_them(self):
        a, b = self.base("A"), self.base("B")
        mapa = self.mapa(a)
        norte, sur = self.folder("bases", "Norte"), self.folder("bases", "Sur")
        informes = self.folder("mapas", "Informes")
        self.move_base(a, norte["id"])
        self.move_base(b, sur["id"])
        self.move_mapa(mapa["id"], informes["id"])
        antes = self.snapshot(mapa["id"])
        terrenos_antes = len(api_bases.terrenos(req(params={"id": a}))["terrenos"])

        out = api_carpetas.remove(req(params={"id": norte["id"]}))
        self.assertEqual(out["trasladados"], 1)
        base_a = api_bases.detail(req(params={"id": a}))["base"]
        self.assertIsNone(base_a["carpeta_id"])
        self.assertEqual(len(api_bases.terrenos(req(params={"id": a}))["terrenos"]), terrenos_antes)
        self.assertEqual(api_bases.detail(req(params={"id": b}))["base"]["carpeta_id"], sur["id"])

        out = api_carpetas.remove(req(params={"id": informes["id"]}))
        self.assertEqual(out, {"eliminada": "Informes", "tipo": "mapas", "trasladados": 1})
        self.assertEqual(self.snapshot(mapa["id"]), antes)
        self.assertIsNone(api_mapas.detail(req(params={"id": mapa["id"]}))["mapa"]["carpeta_id"])

    def test_a_failure_rolls_the_whole_deletion_back(self):
        base_id = self.base()
        carpeta = self.folder("bases", "Frágil")
        self.move_base(base_id, carpeta["id"])

        real = self.conn

        class Failing:
            def execute(self, sql, params=()):
                if sql.startswith("DELETE FROM carpeta"):
                    raise sqlite3.OperationalError("disco lleno")
                return real.execute(sql, params)

        with self.assertRaises(sqlite3.OperationalError):
            repo.delete(Failing(), carpeta["id"])
        self.assertEqual(api_bases.detail(req(params={"id": base_id}))["base"]["carpeta_id"],
                         carpeta["id"])
        self.assertEqual(len(self.folders("bases")["carpetas"]), 1)

    def test_deleting_an_item_leaves_its_folder(self):
        base_id = self.base()
        carpeta = self.folder("bases", "Queda")
        self.move_base(base_id, carpeta["id"])
        api_bases.remove(req(params={"id": base_id}))
        self.assertEqual(self.folders("bases")["carpetas"][0]["conteo"], 0)


class CreationDestinations(FolderCase):
    def preview(self, body, filename):
        return api_importar.preview(req(body=body, headers={"X-Archivo": quote(filename)}))

    def confirm(self, preview, **extra):
        return api_importar.confirm(req(json_body={"token": preview["token"], **extra}))["base"]

    def test_excel_and_csv_imports_land_in_the_chosen_folder(self):
        carpeta = self.folder("bases", "Importadas")
        excel = self.confirm(self.preview(FIXTURE.read_bytes(), "b.xlsx"), carpeta_id=carpeta["id"])
        csv = self.confirm(self.preview(CSV, "b.csv"), carpeta_id=carpeta["id"])
        omitted = self.confirm(self.preview(CSV, "c.csv"))  # old clients send nothing
        self.assertEqual((excel["carpeta_id"], csv["carpeta_id"]), (carpeta["id"], carpeta["id"]))
        self.assertIsNone(omitted["carpeta_id"])

    def test_a_bad_destination_imports_nothing_and_keeps_the_preview_usable(self):
        map_folder = self.folder("mapas", "Mapas")
        preview = self.preview(CSV, "b.csv")
        for carpeta_id, status in ((map_folder["id"], 400), (999, 404), ("1", 400)):
            with self.subTest(carpeta_id=carpeta_id):
                self.assertEqual(self.status_of(
                    lambda c=carpeta_id: self.confirm(preview, carpeta_id=c)), status)
        self.assertEqual(api_bases.listing(req())["bases"], [])
        # The token was not spent on the failed attempts.
        self.assertIsNone(self.confirm(preview, carpeta_id=None)["carpeta_id"])

    def test_every_map_creation_path_honours_its_destination(self):
        a, b = self.base("A"), self.base("B")
        carpeta = self.folder("mapas", "Destino")
        simple = self.mapa(a, carpeta_id=carpeta["id"])
        comparacion = api_mapas.create(req(json_body={
            "nombre": "A vs B", "tipo": "comparacion", "carpeta_id": carpeta["id"],
            "capas": [{"base_id": a, "color": "#111111"}, {"base_id": b, "color": "#222222"}],
        }))["mapa"]
        otro = self.mapa(b, "B")
        combinado = api_mapas.merge(req(json_body={
            "nombre": "Combinado", "mapa_ids": [simple["id"], otro["id"]], "carpeta_id": carpeta["id"],
        }))["mapa"]
        for mapa in (simple, comparacion, combinado):
            self.assertEqual(mapa["carpeta_id"], carpeta["id"])
        self.assertIsNone(otro["carpeta_id"])

    def test_bad_map_destinations_create_nothing(self):
        a, b = self.base("A"), self.base("B")
        base_folder = self.folder("bases", "Bases")
        simple, otro = self.mapa(a), self.mapa(b, "B")
        antes = len(api_mapas.listing(req())["mapas"])
        for carpeta_id, status in ((base_folder["id"], 400), (999, 404), (False, 400)):
            with self.subTest(carpeta_id=carpeta_id):
                self.assertEqual(self.status_of(
                    lambda c=carpeta_id: self.mapa(a, carpeta_id=c)), status)
                self.assertEqual(self.status_of(lambda c=carpeta_id: api_mapas.merge(req(json_body={
                    "nombre": "X", "mapa_ids": [simple["id"], otro["id"]], "carpeta_id": c,
                }))), status)
        self.assertEqual(len(api_mapas.listing(req())["mapas"]), antes)


class ExistingActionsKeepMembership(FolderCase):
    def test_append_rename_refresh_and_saving_a_view(self):
        base_id = self.base("Origen")
        mapa = self.mapa(base_id, "Origen")
        base_folder, map_folder = self.folder("bases", "B"), self.folder("mapas", "M")
        self.move_base(base_id, base_folder["id"])
        self.move_mapa(mapa["id"], map_folder["id"])

        preview = api_importar.preview(req(body=CSV, headers={"X-Archivo": "x.csv"},
                                           query={"base_id": [str(base_id)]}))
        api_importar.append(req(params={"id": base_id}, json_body={"token": preview["token"]}))
        api_bases.rename(req(params={"id": base_id}, json_body={"nombre": "Renombrada"}))
        api_mapas.refresh(req(params={"id": mapa["id"]}))
        api_mapas.update(req(params={"id": mapa["id"]}, json_body={"config": {"basemap": "satelite"}}))
        api_mapas.update(req(params={"id": mapa["id"]}, json_body={"nombre": "Otro nombre"}))

        self.assertEqual(api_bases.detail(req(params={"id": base_id}))["base"]["carpeta_id"],
                         base_folder["id"])
        self.assertEqual(api_mapas.detail(req(params={"id": mapa["id"]}))["mapa"]["carpeta_id"],
                         map_folder["id"])

    def test_deleting_a_source_base_keeps_its_filed_map(self):
        base_id = self.base()
        mapa = self.mapa(base_id)
        carpeta = self.folder("mapas", "Archivo")
        self.move_mapa(mapa["id"], carpeta["id"])
        api_bases.remove(req(params={"id": base_id}))
        guardado = api_mapas.detail(req(params={"id": mapa["id"]}))["mapa"]
        self.assertEqual((guardado["carpeta_id"], guardado["conteo"]), (carpeta["id"], 79))

    def test_a_comparison_across_folders_exports_one_sheet_per_base(self):
        a, b = self.base("Norte"), self.base("Sur")
        self.move_base(a, self.folder("bases", "Cliente 1")["id"])
        self.move_base(b, self.folder("bases", "Cliente 2")["id"])
        mapa = api_mapas.create(req(json_body={
            "nombre": "Clientes", "tipo": "comparacion",
            "capas": [{"base_id": a, "color": "#111111"}, {"base_id": b, "color": "#222222"}],
        }))["mapa"]
        payload, _ = api_exportar.export(req(json_body={"mapa_id": mapa["id"]}))
        self.assertEqual(load_workbook(io.BytesIO(payload)).sheetnames, ["Norte", "Sur"])

    def test_listings_always_carry_carpeta_id(self):
        base_id = self.base()
        self.mapa(base_id)
        self.assertIn("carpeta_id", api_bases.listing(req())["bases"][0])
        self.assertIn("carpeta_id", api_mapas.listing(req())["mapas"][0])
        self.assertIn("carpeta_id", json.loads(json.dumps(api_bases.detail(req(params={"id": base_id})))) ["base"])
