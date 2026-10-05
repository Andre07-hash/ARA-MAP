"""API handler tests, driven directly rather than over HTTP."""

from __future__ import annotations

import json
import unittest

from server.api import bases as api_bases
from server.api import exportar as api_exportar
from server.api import importar as api_importar
from server.api import mapas as api_mapas
from server.router import Request
from server.web_util import ApiError
from tests.support import FIXTURE, TempDatabase


def req(*, body=b"", params=None, query=None, headers=None, json_body=None):
    if json_body is not None:
        body = json.dumps(json_body).encode()
    return Request("POST", "/x", query or {}, params or {}, body, headers or {})


class BasesEndpoints(TempDatabase):
    def test_listing_starts_empty(self):
        self.assertEqual(api_bases.listing(req())["bases"], [])

    def test_detail_and_terrains_after_an_import(self):
        base_id, _ = self.load_fixture("Septiembre")
        detail = api_bases.detail(req(params={"id": base_id}))["base"]
        self.assertEqual(detail["nombre"], "Septiembre")
        self.assertEqual(detail["conteo"], 79)

        rows = api_bases.terrenos(req(params={"id": base_id}))["terrenos"]
        self.assertEqual(len(rows), 79)
        self.assertEqual(sum(1 for r in rows if r["ubicado"]), 39)

    def test_unknown_base_is_a_404(self):
        for handler in (api_bases.detail, api_bases.terrenos, api_bases.remove):
            with self.subTest(handler=handler.__name__):
                with self.assertRaises(ApiError) as caught:
                    handler(req(params={"id": 999}))
                self.assertEqual(caught.exception.status, 404)

    def test_rename_requires_a_name(self):
        base_id, _ = self.load_fixture()
        with self.assertRaises(ApiError):
            api_bases.rename(req(params={"id": base_id}, json_body={"nombre": ""}))

    def test_rename_trims_and_persists(self):
        base_id, _ = self.load_fixture()
        out = api_bases.rename(req(params={"id": base_id}, json_body={"nombre": "  Octubre  "}))
        self.assertEqual(out["base"]["nombre"], "Octubre")

    def test_remove_reports_what_it_deleted(self):
        base_id, _ = self.load_fixture("Septiembre")
        out = api_bases.remove(req(params={"id": base_id}))
        self.assertEqual(out, {"eliminada": "Septiembre", "terrenos": 79})

    def test_adding_a_terrain_by_hand_derives_the_missing_area_unit(self):
        base_id, _ = self.load_fixture()
        api_bases.add_terreno(req(params={"id": base_id}, json_body={
            "terreno": "Nuevo predio", "estado": "Jalisco", "municipio": "Tala",
            "superficie_m2": 25000, "lat": 20.6, "lon": -103.7,
        }))
        rows = api_bases.terrenos(req(params={"id": base_id}))["terrenos"]
        added = next(r for r in rows if r["terreno"] == "Nuevo predio")
        self.assertEqual(added["superficie_ha"], 2.5)
        self.assertTrue(added["ubicado"])

    def test_a_manual_terrain_needs_a_name(self):
        base_id, _ = self.load_fixture()
        with self.assertRaises(ApiError):
            api_bases.add_terreno(req(params={"id": base_id}, json_body={"estado": "Jalisco"}))

    def test_a_non_numeric_field_is_rejected_with_its_name(self):
        base_id, _ = self.load_fixture()
        with self.assertRaises(ApiError) as caught:
            api_bases.add_terreno(req(params={"id": base_id},
                                      json_body={"terreno": "X", "superficie_m2": "mucho"}))
        self.assertIn("superficie_m2", caught.exception.mensaje)


class ImportEndpoints(TempDatabase):
    def upload(self, base_id=None, filename="Base%20Terrenos%2009.26.xlsx"):
        return api_importar.preview(req(
            body=FIXTURE.read_bytes(),
            headers={"X-Archivo": filename},
            query={"base_id": [str(base_id)]} if base_id else {},
        ))

    def test_preview_reports_the_file_without_writing_anything(self):
        preview = self.upload()
        self.assertEqual(preview["conteo"], 79)
        self.assertEqual(preview["ubicados"], 39)
        self.assertEqual(preview["sin_ubicacion"], 40)
        self.assertEqual(preview["hoja"], "Registro Análisis")
        self.assertEqual(preview["nombre_sugerido"], "Base Terrenos 09.26")
        self.assertEqual(api_bases.listing(req())["bases"], [])

    def test_preview_surfaces_the_price_typo(self):
        self.assertEqual(self.upload()["resumen"]["PRECIO_INCONSISTENTE"], 1)

    def test_an_empty_upload_is_rejected(self):
        with self.assertRaises(ApiError):
            api_importar.preview(req(body=b""))

    def test_a_numbers_file_gets_a_useful_message(self):
        with self.assertRaises(ApiError) as caught:
            api_importar.preview(req(body=b"x", headers={"X-Archivo": "base.numbers"}))
        self.assertIn("Numbers", caught.exception.mensaje)

    def test_a_file_that_is_not_a_workbook_is_reported(self):
        with self.assertRaises(ApiError):
            api_importar.preview(req(body=b"no soy excel",
                                     headers={"X-Archivo": "base.xlsx"}))

    def test_confirm_writes_the_previewed_rows(self):
        preview = self.upload()
        out = api_importar.confirm(req(json_body={"token": preview["token"],
                                                  "nombre": "Septiembre"}))
        self.assertEqual(out["base"]["nombre"], "Septiembre")
        self.assertEqual(out["base"]["conteo"], 79)

    def test_a_token_cannot_be_reused(self):
        preview = self.upload()
        api_importar.confirm(req(json_body={"token": preview["token"], "nombre": "A"}))
        with self.assertRaises(ApiError) as caught:
            api_importar.confirm(req(json_body={"token": preview["token"], "nombre": "B"}))
        self.assertEqual(caught.exception.status, 410)

    def test_a_colliding_name_is_suffixed_rather_than_overwritten(self):
        api_importar.confirm(req(json_body={"token": self.upload()["token"], "nombre": "Base"}))
        out = api_importar.confirm(req(json_body={"token": self.upload()["token"],
                                                  "nombre": "Base"}))
        self.assertEqual(out["base"]["nombre"], "Base (2)")

    def test_appending_the_same_file_classifies_everything_as_duplicate(self):
        base_id, _ = self.load_fixture()
        preview = self.upload(base_id=base_id)
        self.assertEqual(preview["clasificacion"]["nuevas"], 0)
        self.assertEqual(preview["clasificacion"]["duplicadas"], 79)
        # Identical content must never be reported as a change.
        self.assertEqual(preview["clasificacion"]["conflictos"], 0)

        out = api_importar.append(req(params={"id": base_id},
                                      json_body={"token": preview["token"]}))
        self.assertEqual(out["agregados"], 0)
        self.assertEqual(out["omitidos"], 79)
        self.assertEqual(out["base"]["conteo"], 79)

    def test_appending_to_an_empty_base_adds_the_rows(self):
        from server.repo import bases as repo_bases
        base_id = repo_bases.create(self.conn, "Vacía", None, None)
        preview = self.upload(base_id=base_id)
        self.assertEqual(preview["clasificacion"]["nuevas"], 78)  # one true duplicate
        out = api_importar.append(req(params={"id": base_id},
                                      json_body={"token": preview["token"]}))
        self.assertEqual(out["agregados"], 78)

    def test_previewing_against_a_missing_base_is_a_404(self):
        with self.assertRaises(ApiError) as caught:
            self.upload(base_id=999)
        self.assertEqual(caught.exception.status, 404)


class MapEndpoints(TempDatabase):
    def setUp(self):
        super().setUp()
        from server.repo import bases as repo_bases
        self.a = repo_bases.create(self.conn, "Agosto", None, None)
        self.b = repo_bases.create(self.conn, "Septiembre", None, None)
        self.c = repo_bases.create(self.conn, "Octubre", None, None)
        self.d = repo_bases.create(self.conn, "Noviembre", None, None)

    def capa(self, base_id, color="#2a78d6"):
        return {"base_id": base_id, "color": color}

    def test_creating_and_reading_a_comparison(self):
        out = api_mapas.create(req(json_body={
            "nombre": "Agosto vs Septiembre", "tipo": "comparacion",
            "capas": [self.capa(self.a), self.capa(self.b, "#eb6834")],
        }))
        mapa_id = out["mapa"]["id"]
        stored = api_mapas.detail(req(params={"id": mapa_id}))["mapa"]
        self.assertEqual(len(stored["capas"]), 2)
        self.assertEqual(api_mapas.listing(req())["mapas"][0]["nombre"], "Agosto vs Septiembre")

    def test_four_layers_are_allowed(self):
        """Four is within the range where colours stay distinguishable."""
        out = api_mapas.create(req(json_body={
            "nombre": "Cuatro", "tipo": "comparacion",
            "capas": [self.capa(self.a), self.capa(self.b),
                      self.capa(self.c), self.capa(self.d)],
        }))
        self.assertEqual(len(out["mapa"]["capas"]), 4)

    def test_a_simple_map_holds_exactly_one_base(self):
        with self.assertRaises(ApiError):
            api_mapas.create(req(json_body={
                "nombre": "M", "tipo": "simple",
                "capas": [self.capa(self.a), self.capa(self.b)]}))

    def test_a_base_cannot_appear_twice(self):
        with self.assertRaises(ApiError):
            api_mapas.create(req(json_body={
                "nombre": "M", "tipo": "comparacion",
                "capas": [self.capa(self.a), self.capa(self.a)]}))

    def test_an_unknown_type_is_refused(self):
        with self.assertRaises(ApiError):
            api_mapas.create(req(json_body={
                "nombre": "M", "tipo": "otro", "capas": [self.capa(self.a)]}))

    def test_a_map_needs_at_least_one_base(self):
        with self.assertRaises(ApiError):
            api_mapas.create(req(json_body={"nombre": "M", "tipo": "simple", "capas": []}))

    def test_a_layer_pointing_at_a_missing_base_is_a_404(self):
        with self.assertRaises(ApiError) as caught:
            api_mapas.create(req(json_body={
                "nombre": "M", "tipo": "simple", "capas": [self.capa(999)]}))
        self.assertEqual(caught.exception.status, 404)

    def test_map_terrains_carry_their_layer_colour(self):
        base_id, _ = self.load_fixture("Con datos")
        out = api_mapas.create(req(json_body={
            "nombre": "M", "tipo": "simple", "capas": [self.capa(base_id, "#1baf7a")]}))
        payload = api_mapas.terrenos(req(params={"id": out["mapa"]["id"]}))
        self.assertEqual(len(payload["terrenos"]), 79)
        self.assertTrue(all(t["color"] == "#1baf7a" for t in payload["terrenos"]))

    def test_update_and_delete(self):
        out = api_mapas.create(req(json_body={
            "nombre": "Antes", "tipo": "simple", "capas": [self.capa(self.a)]}))
        mapa_id = out["mapa"]["id"]
        renamed = api_mapas.update(req(params={"id": mapa_id}, json_body={"nombre": "Después"}))
        self.assertEqual(renamed["mapa"]["nombre"], "Después")
        self.assertEqual(api_mapas.remove(req(params={"id": mapa_id})),
                         {"eliminado": "Después"})

    def test_unknown_map_is_a_404(self):
        for handler in (api_mapas.detail, api_mapas.terrenos,
                        api_mapas.remove, api_mapas.update):
            with self.subTest(handler=handler.__name__):
                with self.assertRaises(ApiError) as caught:
                    handler(req(params={"id": 999}, json_body={"nombre": "x"}))
                self.assertEqual(caught.exception.status, 404)


class ExportEndpoint(TempDatabase):
    def test_exports_a_readable_workbook(self):
        import io

        from openpyxl import load_workbook

        base_id, _ = self.load_fixture("Septiembre")
        payload, filename = api_exportar.export(
            req(json_body={"base_id": base_id, "nombre": "Septiembre"}))
        self.assertEqual(filename, "Septiembre.xlsx")

        sheet = load_workbook(io.BytesIO(payload)).active
        self.assertEqual(sheet.title, "Registro Análisis")
        self.assertEqual(sheet.max_row, 80)  # header + 79
        self.assertEqual(sheet.cell(row=1, column=2).value, "Terreno")

    def test_only_the_requested_rows_are_exported(self):
        import io

        from openpyxl import load_workbook

        base_id, _ = self.load_fixture()
        rows = api_bases.terrenos(req(params={"id": base_id}))["terrenos"]
        wanted = [r["id"] for r in rows[:5]]
        payload, _ = api_exportar.export(
            req(json_body={"base_id": base_id, "ids": wanted}))
        self.assertEqual(load_workbook(io.BytesIO(payload)).active.max_row, 6)

    def test_exporting_nothing_is_an_error(self):
        base_id, _ = self.load_fixture()
        with self.assertRaises(ApiError):
            api_exportar.export(req(json_body={"base_id": base_id, "ids": []}))

    def test_a_filename_cannot_escape_its_directory(self):
        base_id, _ = self.load_fixture()
        _, filename = api_exportar.export(
            req(json_body={"base_id": base_id, "nombre": "../../etc/passwd"}))
        self.assertNotIn("/", filename)


if __name__ == "__main__":
    unittest.main()
