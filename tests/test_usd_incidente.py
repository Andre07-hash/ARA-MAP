"""The October 1, 2026 incident: the real terrain workbook, in US dollars.

reports/urgent-usd-import-2026-10-01/DEVELOPER_HANDOFF.md. The supervisor
confirmed the source prices are US dollars; the workbook's "[$$-409]" format
does not say so by itself, so the import asks once and the answer is USD.
Everything runs through the real handlers on a throwaway database.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from server.api import bases as api_bases
from server.api import exportar as api_exportar
from server.api import mapas as api_mapas
from server.repo import formatos as repo_formatos
from server.web_util import ApiError
from tests.support import FIXTURE
from tests.test_api import req
from tests.test_asistente_api import AsistenteCase

# Production format id 4, as reconstructed in the handoff's evidence: saved by
# the version that refused dollars, it keeps both price columns as extra data.
FORMATO_4 = [
    ["id", 1, "id_origen"], ["terreno", 1, "terreno"], ["estado", 1, "estado"],
    ["municipio", 1, "municipio"], ["direccion", 1, "direccion"],
    ["superficie m2", 1, "superficie_m2"], ["superficie ha", 1, "superficie_ha"],
    ["afectaciones %", 1, "afectaciones_pct"], ["afectaciones m2", 1, "afectaciones_m2"],
    ["asking price", 1, "extra"], ["asking $/m2", 1, "extra"], ["x", 1, "lat"], ["y", 1, "lon"],
]


class ArchivoReal(AsistenteCase):
    moneda = None   # the currency is the point here: answered explicitly below

    def analizar_real(self, contenido=None, nombre="Base Terrenos 09.26 copy.xlsx"):
        return self.analizar(nombre, contenido if contenido is not None else FIXTURE.read_bytes())

    def responder(self, r, **elecciones):
        """Answer open questions: moneda=USD, and «Sí» to restoring a price column."""
        respuestas = []
        for pregunta in r["preguntas"]:
            if pregunta["id"] == "moneda":
                codigo = elecciones["moneda"]
                opcion = next(o["indice"] for o in pregunta["opciones"] if f"({codigo})" in o["etiqueta"])
            else:
                opcion = next(o["indice"] for o in pregunta["opciones"] if o["etiqueta"].startswith("Sí"))
            respuestas.append({"pregunta": pregunta["id"], "opcion": opcion})
        return self.preparar(r, respuestas)

    def terrenos(self, base_id):
        return {t["terreno"]: t for t in api_bases.terrenos(req(params={"id": base_id}))["terrenos"]}

    # -- acceptance 1: preview -----------------------------------------------

    def test_the_original_workbook_asks_its_currency_once(self):
        r = self.analizar_real()
        self.assertEqual(r["estado"], "preguntas")
        self.assertEqual([p["id"] for p in r["preguntas"]], ["moneda"])
        self.assertIsNone(r["preguntas"][0]["sugerida"])   # never guessed from magnitudes

    def test_preview_counts_and_currency(self):
        r = self.responder(self.analizar_real(), moneda="USD")
        vista = r["vista_previa"]
        self.assertEqual((vista["conteo"], vista["ubicados"], vista["con_precio"], vista["con_precio_m2"]),
                         (79, 39, 59, 60))
        self.assertEqual((vista["moneda"], vista["moneda_origen"]), ("USD", "usuario"))
        self.assertEqual(vista["monedas_no_admitidas"], [])
        self.assertEqual(vista["hoja"], "Registro Análisis")

    # -- acceptance 2: stored, reopened ---------------------------------------

    def test_confirmed_values_are_the_source_values_in_usd(self):
        r = self.responder(self.analizar_real(), moneda="USD")
        base = self.confirmar(r, nombre="Base Terrenos 09.26")["base"]
        self.assertEqual((base["conteo"], base["ubicados"]), (79, 39))
        terrenos = self.terrenos(base["id"])
        marcenas, dorado = terrenos["Marceñas"], terrenos["El Dorado"]
        self.assertEqual((marcenas["asking_price"], marcenas["asking_m2"], marcenas["moneda"]),
                         (388689722.0, 600.0, "USD"))
        self.assertEqual((dorado["asking_m2"], dorado["moneda"]), (122.5, "USD"))
        # J2 disagrees with area x unit price; the source value is kept and flagged.
        self.assertIn("PRECIO_INCONSISTENTE", {i["codigo"] for i in marcenas["incidencias"]})
        self.assertTrue(any("USD" in i["mensaje"] for i in marcenas["incidencias"]))
        con_precio = [t for t in terrenos.values() if t["asking_price"] is not None or t["asking_m2"] is not None]
        self.assertEqual({t["moneda"] for t in con_precio}, {"USD"})
        sin_coordenadas = [t for t in terrenos.values() if not t["ubicado"]]
        self.assertEqual(len(sin_coordenadas), 40)
        self.assertTrue(all(t["lat"] is None or t["lon"] is None or t["ubicacion"] == "invalida"
                            for t in sin_coordenadas))

    # -- acceptance 3: the stale saved format ---------------------------------

    def test_format_4_cannot_silently_drop_the_prices(self):
        firma = [[c, o] for c, o, _ in FORMATO_4]
        repo_formatos.recordar(self.conn, "Formato de «Base Terrenos 09.26 copy»", firma,
                               {"asignaciones": FORMATO_4, "decimal": "dot", "hoja": "CSV", "formato": "csv"}, 1)
        r = self.analizar_real()
        self.assertEqual(r["estado"], "preguntas")
        self.assertEqual(sorted(p["id"] for p in r["preguntas"])[-1], "moneda")
        self.assertEqual(len(r["preguntas"]), 3)
        self.assertTrue(all("dato adicional" in p["detalle"] for p in r["preguntas"] if p["id"] != "moneda"))

        r = self.responder(r, moneda="USD")
        vista = r["vista_previa"]
        self.assertEqual((vista["conteo"], vista["ubicados"], vista["con_precio"], vista["con_precio_m2"],
                          vista["moneda"]), (79, 39, 59, 60, "USD"))
        salida = self.confirmar(r, nombre="Base Terrenos 09.26", recordar_formato=True)
        self.assertEqual(salida["formato"]["accion"], "nueva_version")
        # Superseded, not deleted; the new version remembers the prices and USD.
        formatos = {f["id"]: f for f in repo_formatos.listing(self.conn)}
        viejo = next(f for f in formatos.values() if f["version"] == 1)
        nuevo = formatos[salida["formato"]["id"]]
        self.assertEqual(viejo["reemplazado_por"], nuevo["id"])
        self.assertEqual(nuevo["moneda"], "USD")
        self.assertIn(["asking price", 1, "asking_price"], nuevo["asignaciones"])

        # The next upload of the same workbook goes straight to a USD preview.
        otra = self.analizar_real()
        self.assertEqual(otra["estado"], "vista_previa")
        self.assertEqual((otra["vista_previa"]["moneda"], otra["vista_previa"]["moneda_origen"],
                          otra["vista_previa"]["con_precio"]), ("USD", "formato", 59))

    # -- acceptance 5: maps, comparison, export and re-import ----------------

    def importar_mxn(self):
        wb = Workbook()
        wb.active.title = "Registro Análisis"
        wb.active.append(["Terreno", "Estado", "Municipio", "Superficie m2", "Asking Price", "Asking $/m2",
                          "X", "Y"])
        wb.active.append(["FICTICIO Pesos", "Jalisco", "Zapopan", 1000, 600000, 600, 20.7, -103.4])
        for celda in ("E2", "F2"):
            wb.active[celda].number_format = '"MXN "#,##0'
        buffer = io.BytesIO()
        wb.save(buffer)
        r = self.analizar("ficticia_mxn.xlsx", buffer.getvalue())
        self.assertEqual(r["vista_previa"]["moneda"], "MXN")
        return self.confirmar(r, nombre="Ficticia MXN")["base"]

    def test_save_compare_export_and_reimport_keep_currency_and_cents(self):
        usd = self.confirmar(self.responder(self.analizar_real(), moneda="USD"), nombre="Real USD")["base"]
        mxn = self.importar_mxn()
        mapa_usd = api_mapas.create(req(json_body={"nombre": "USD", "tipo": "simple",
                                                   "capas": [{"base_id": usd["id"], "color": "#111111"}]}))["mapa"]
        mapa_mxn = api_mapas.create(req(json_body={"nombre": "MXN", "tipo": "simple",
                                                   "capas": [{"base_id": mxn["id"], "color": "#222222"}]}))["mapa"]
        congelados = api_mapas.terrenos(req(params={"id": mapa_usd["id"]}))["terrenos"]
        dorado = next(t for t in congelados if t["terreno"] == "El Dorado")
        self.assertEqual((dorado["asking_m2"], dorado["moneda"]), (122.5, "USD"))

        comparacion = api_mapas.merge(req(json_body={"nombre": "Comparación", "mapa_ids": [mapa_usd["id"],
                                                     mapa_mxn["id"]]}))["mapa"]
        filas = api_mapas.terrenos(req(params={"id": comparacion["id"]}))["terrenos"]
        self.assertEqual({(t["capa"], t["moneda"]) for t in filas if t["asking_price"] is not None},
                         {(0, "USD"), (1, "MXN")})
        self.assertEqual(sum(t["ubicado"] for t in filas), 40)   # geography still compares

        contenido, _ = api_exportar.export(req(json_body={"mapa_id": comparacion["id"], "nombre": "c"}))
        libro = load_workbook(io.BytesIO(contenido))
        self.assertEqual(len(libro.worksheets), 2)   # one sheet per source
        hoja = libro.worksheets[0]
        encabezados = [c.value for c in hoja[1]]
        fila = next(r for r in hoja.iter_rows(min_row=2) if r[1].value == "El Dorado")
        celda = fila[encabezados.index("Asking $/m2")]
        self.assertEqual((celda.value, fila[encabezados.index("Moneda")].value), (122.5, "USD"))
        self.assertIn('"USD "', celda.number_format)
        self.assertIn("0.00", celda.number_format)
        pesos = next(r for r in libro.worksheets[1].iter_rows(min_row=2) if r[1].value == "FICTICIO Pesos")
        self.assertEqual(pesos[encabezados.index("Moneda")].value, "MXN")

        # Re-importing the USD sheet needs no question: the file now says USD.
        solo_usd = io.BytesIO()
        libro.remove(libro.worksheets[1])
        libro.save(solo_usd)
        r = self.analizar("exportado.xlsx", solo_usd.getvalue())
        self.assertEqual(r["estado"], "vista_previa", r.get("preguntas"))
        vista = r["vista_previa"]
        self.assertEqual((vista["conteo"], vista["con_precio"], vista["con_precio_m2"], vista["moneda"],
                          vista["moneda_origen"]), (79, 59, 60, "USD", "archivo"))
        base = self.confirmar(r, nombre="Reimportada")["base"]
        self.assertEqual(self.terrenos(base["id"])["El Dorado"]["asking_m2"], 122.5)

    def test_a_legacy_row_without_currency_is_exported_as_unconfirmed(self):
        base_id, _ = self.load_fixture("Legado")
        contenido, _ = api_exportar.export(req(json_body={"base_id": base_id, "nombre": "l"}))
        hoja = load_workbook(io.BytesIO(contenido)).worksheets[0]
        encabezados = [c.value for c in hoja[1]]
        monedas = {r[encabezados.index("Moneda")].value for r in hoja.iter_rows(min_row=2)
                   if r[encabezados.index("Asking Price")].value is not None}
        self.assertEqual(monedas, {api_exportar.SIN_CONFIRMAR})
        # ...and re-importing it cannot quietly give those amounts a currency.
        r = self.analizar("legado.xlsx", contenido)
        self.assertNotIn("vista_previa", r)

    def test_a_manual_price_needs_its_currency(self):
        base_id, _ = self.load_fixture("Manual")
        with self.assertRaises(ApiError):
            api_bases.add_terreno(req(params={"id": base_id}, json_body={"terreno": "A", "asking_price": 5}))
        api_bases.add_terreno(req(params={"id": base_id},
                                  json_body={"terreno": "B", "asking_price": 5, "moneda": "USD"}))
        self.assertEqual(self.terrenos(base_id)["B"]["moneda"], "USD")
