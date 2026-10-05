"""CSV through the import API, staging, maps and export. Fictional data only.

TempDatabase isolates SQLite; the cloud database variables are also cleared so
nothing here can reach a shared Postgres workspace.
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from contextlib import contextmanager
from dataclasses import fields
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import Workbook, load_workbook

from server.api import bases as api_bases
from server.api import exportar as api_exportar
from server.api import importar as api_importar
from server.api import mapas as api_mapas
from server.csv_importer import read_csv_bytes
from server.importer import SHEET_NAME, TerrainRecord, read_workbook
from server.staging import Staging
from server.validation import validate_all
from server.web_util import ApiError
from tests.support import FIXTURE, TempDatabase
from tests.test_api import req

# One canonical set of values, written out both as CSV text and as typed Excel
# cells, so the two readers can be compared value by value.
COLUMNS = ["ID", "Terreno", "Estado", "Municipio", "Dirección", "Superficie m2", "Superficie Ha",
           "Afectaciones %", "Afectaciones m2", "Asking Price", "Asking $/m2", "X", "Y"]
CANONICAL = [
    [1, "Predio Prueba Norte", "Estado de México", "Tecámac", "Av. Central 120, Colonia Centro",
     25000, 2.5, 0.10, 2500, 12500000, 500, 19.713, -98.968],
    [2, "Predio Prueba Sur", "Jalisco", "Tala", "Camino Rural 8",
     40000, 4, 0, 0, 12000000, 300, 20.653, -103.701],
    [3, "Predio Sin Mapa", "Sonora", "Hermosillo", None,
     10000, 1, None, None, 3000000, 300, None, None],
]
CSV_TEXT = (
    ",".join(COLUMNS) + "\n"
    '1,Predio Prueba Norte,Estado de México,Tecámac,"Av. Central 120, Colonia Centro",'
    '25000,2.5,10%,2500,"12,500,000",500,19.713,-98.968\n'
    "2,Predio Prueba Sur,Jalisco,Tala,Camino Rural 8,40000,4,0,0,12000000,300,20.653,-103.701\n"
    "3,Predio Sin Mapa,Sonora,Hermosillo,,10000,1,,,3000000,300,,\n"
)
CSV_BYTES = CSV_TEXT.encode()


def excel_bytes(rows=CANONICAL) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class NoCloud(TempDatabase):
    def setUp(self):
        self._env = patch.dict(os.environ, {"ARA_MAP_DATABASE_URL": "", "DATABASE_URL": ""})
        self._env.start()
        super().setUp()

    def tearDown(self):
        super().tearDown()
        self._env.stop()

    def preview(self, body=CSV_BYTES, filename="terrenos.csv", **query):
        return api_importar.preview(req(
            body=body, headers={"X-Archivo": quote(filename)},
            query={k: [str(v)] for k, v in query.items()},
        ))

    def confirm(self, preview, nombre="CSV"):
        return api_importar.confirm(req(json_body={"token": preview["token"], "nombre": nombre}))["base"]

    def terrenos(self, base_id):
        return api_bases.terrenos(req(params={"id": base_id}))["terrenos"]


class Dispatch(NoCloud):
    def test_a_real_csv_previews(self):
        preview = self.preview(filename="Base Acentuada ñ.CSV")
        self.assertEqual((preview["formato"], preview["csv_decimal"], preview["hoja"]), ("csv", "dot", "CSV"))
        self.assertEqual((preview["conteo"], preview["ubicados"], preview["sin_ubicacion"]), (3, 2, 1))
        self.assertEqual(preview["nombre_sugerido"], "Base Acentuada ñ")
        self.assertEqual(preview["filas_con_datos"], 3)

    def test_excel_preview_reports_its_format_and_no_decimal(self):
        preview = self.preview(body=excel_bytes(), filename="base.xlsm")
        self.assertEqual(preview["formato"], "xlsm")
        self.assertNotIn("csv_decimal", preview)
        self.assertEqual(preview["hoja"], SHEET_NAME)

    def test_excel_ignores_the_csv_parameter(self):
        preview = self.preview(body=excel_bytes(), filename="base.xlsx", csv_decimal="bogus")
        self.assertEqual(preview["conteo"], 3)

    def test_unknown_decimal_mode_is_rejected_for_csv(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(csv_decimal="space")
        self.assertIn("csv_decimal", caught.exception.mensaje)

    def test_comma_mode_changes_the_reading(self):
        text = b"Terreno;X;Y;Superficie m2\nPredio A;19,713;-98,968;1.234,5\n"
        dot = self.preview(body=text)
        self.assertEqual(dot["ubicados"], 0)
        comma = self.preview(body=text, csv_decimal="comma")
        self.assertEqual((comma["csv_decimal"], comma["ubicados"]), ("comma", 1))
        self.assertNotEqual(dot["token"], comma["token"])

    def test_excel_renamed_csv_is_rejected(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(body=excel_bytes())
        self.assertIn("libro de Excel", caught.exception.mensaje)

    def test_csv_renamed_xlsx_gets_the_excel_error(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(filename="base.xlsx")
        self.assertIn("Excel", caught.exception.mensaje)

    def test_unsupported_suffix(self):
        for name in ("base.numbers", "base.txt", "base.csv.bak", "base"):
            with self.subTest(name=name), self.assertRaises(ApiError) as caught:
                self.preview(filename=name)
            self.assertIn(".xlsx, .xlsm o .csv", caught.exception.mensaje)

    def test_non_utf8_fails_usefully(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(body="Terreno\nTecámac\n".encode("cp1252"))
        self.assertIn("CSV UTF-8", caught.exception.mensaje)

    def test_the_size_limit_applies_to_csv(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(body=b"Terreno\n" + b"x" * api_importar.MAX_UPLOAD)
        self.assertEqual(caught.exception.status, 413)

    def test_all_rejected_is_an_error_and_writes_nothing(self):
        with self.assertRaises(ApiError) as caught:
            self.preview(body=b"Terreno,Estado\n,Jalisco\n")
        self.assertIn("CSV no contiene", caught.exception.mensaje)
        self.assertEqual(api_bases.listing(req())["bases"], [])

    def test_malformed_file_is_an_error_and_writes_nothing(self):
        with self.assertRaises(ApiError):
            self.preview(body=b"Terreno,Estado\nA,B,C\n")
        self.assertEqual(api_bases.listing(req())["bases"], [])

    def test_duplicate_aliases_are_reported_as_such(self):
        preview = self.preview(body=b"Terreno,Nombre,Vendedor\nA,B,C\n")
        self.assertEqual(preview["columnas_no_reconocidas"], ["Nombre", "Vendedor"])
        self.assertEqual(preview["columnas_duplicadas"], ["Nombre"])

    def test_findings_carry_the_line_and_original_text(self):
        body = b"Terreno,Estado,Municipio,X,Y\nA,Jalisco,Tala,\"19,4326\",-99\n"
        hallazgo = self.preview(body=body)["hallazgos"][0]
        self.assertEqual(hallazgo["fila"], 2)
        texto = " ".join(i["mensaje"] for i in hallazgo["incidencias"])
        self.assertIn("«19,4326»", texto)
        self.assertIn("punto decimal", texto)


class PreviewAndConfirm(NoCloud):
    def test_preview_writes_nothing(self):
        self.preview()
        self.assertEqual(api_bases.listing(req())["bases"], [])
        self.assertEqual(api_mapas.listing(req())["mapas"], [])

    def test_confirm_creates_exactly_the_previewed_records(self):
        base = self.confirm(self.preview(filename="Mi CSV.csv"), "Desde CSV")
        self.assertEqual((base["nombre"], base["conteo"], base["ubicados"]), ("Desde CSV", 3, 2))
        self.assertEqual((base["archivo_origen"], base["hoja"]), ("Mi CSV.csv", "CSV"))

        rows = {r["terreno"]: r for r in self.terrenos(base["id"])}
        norte = rows["Predio Prueba Norte"]
        self.assertEqual((norte["lat"], norte["lon"]), (19.713, -98.968))
        self.assertAlmostEqual(norte["afectaciones_pct"], 0.10)
        self.assertEqual(norte["asking_price"], 12_500_000)
        self.assertEqual(norte["direccion"], "Av. Central 120, Colonia Centro")
        self.assertFalse(rows["Predio Sin Mapa"]["ubicado"])

    def test_a_token_cannot_be_reused(self):
        preview = self.preview()
        self.confirm(preview)
        with self.assertRaises(ApiError) as caught:
            self.confirm(preview)
        self.assertEqual(caught.exception.status, 410)

    def test_an_expired_token_fails(self):
        preview = self.preview()
        with patch("server.staging.time.time", return_value=10**12), \
                self.assertRaises(ApiError) as caught:
            self.confirm(preview)
        self.assertEqual(caught.exception.status, 410)

    def test_cancelling_leaves_nothing(self):
        self.preview()  # the dialog is closed without confirming
        self.assertEqual(api_bases.listing(req())["bases"], [])


class Append(NoCloud):
    def test_csv_into_an_excel_created_base(self):
        base_id, _ = self.load_fixture()
        preview = self.preview(base_id=base_id)
        self.assertEqual(preview["clasificacion"]["nuevas"], 3)
        out = api_importar.append(req(params={"id": base_id}, json_body={"token": preview["token"]}))
        self.assertEqual((out["agregados"], out["base"]["conteo"]), (3, 82))

    def test_excel_into_a_csv_created_base_is_all_duplicates(self):
        base = self.confirm(self.preview())
        preview = self.preview(body=excel_bytes(), filename="misma.xlsx", base_id=base["id"])
        clasificacion = preview["clasificacion"]
        self.assertEqual((clasificacion["nuevas"], clasificacion["duplicadas"], clasificacion["conflictos"]),
                         (0, 3, 0))

    def test_conflicts_keep_or_update_as_chosen(self):
        base = self.confirm(self.preview())
        changed = CSV_TEXT.replace("25000,2.5,10%,2500,\"12,500,000\",500",
                                   "25000,2.5,10%,2500,\"13,000,000\",520")
        changed = changed.replace("40000,4,0,0,12000000,300", "40000,4,0,0,12800000,320")
        preview = self.preview(body=changed.encode(), base_id=base["id"])
        detalle = preview["clasificacion"]["detalle"]
        self.assertEqual(preview["clasificacion"]["conflictos"], 2)
        indices = {d["terreno"]: d["indice"] for d in detalle}

        out = api_importar.append(req(params={"id": base["id"]}, json_body={
            "token": preview["token"],
            "resoluciones": {str(indices["Predio Prueba Norte"]): "actualizar",
                             str(indices["Predio Prueba Sur"]): "omitir"},
        }))
        self.assertEqual((out["agregados"], out["actualizados"], out["omitidos"]), (0, 1, 2))
        rows = {r["terreno"]: r for r in self.terrenos(base["id"])}
        self.assertEqual(rows["Predio Prueba Norte"]["asking_price"], 13_000_000)
        self.assertEqual(rows["Predio Prueba Sur"]["asking_price"], 12_000_000)


class MapsAndExport(NoCloud):
    def test_csv_base_in_a_saved_comparison_and_its_export(self):
        excel_id, _ = self.load_fixture("Septiembre")
        csv_base = self.confirm(self.preview(), "Octubre CSV")

        mapa = api_mapas.create(req(json_body={
            "nombre": "Sep vs Oct", "tipo": "comparacion",
            "capas": [{"base_id": excel_id, "color": "#2a78d6"},
                      {"base_id": csv_base["id"], "color": "#eb6834"}],
        }))["mapa"]
        reabierto = api_mapas.detail(req(params={"id": mapa["id"]}))["mapa"]
        self.assertEqual(len(reabierto["capas"]), 2)

        terrenos = api_mapas.terrenos(req(params={"id": mapa["id"]}))["terrenos"]
        del_csv = [t for t in terrenos if t["color"] == "#eb6834"]
        self.assertEqual(len(del_csv), 3)  # the unlocated one stays in the data
        self.assertEqual(sum(1 for t in del_csv if t["ubicado"]), 2)

        payload, _ = api_exportar.export(req(json_body={"mapa_id": mapa["id"], "nombre": "cmp"}))
        workbook = load_workbook(io.BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ["Septiembre", "Octubre CSV"])
        self.assertEqual(workbook["Octubre CSV"].max_row, 4)


class ExcelEquivalence(unittest.TestCase):
    """The same canonical values, read from Excel and from CSV, must agree."""

    SOURCE_ONLY = {"fila", "notes", "extra"}

    def test_same_values_same_findings(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "canonica.xlsx"
            path.write_bytes(excel_bytes())
            excel = read_workbook(path)
        csv = read_csv_bytes(CSV_BYTES)

        self.assertEqual(len(excel.records), len(csv.records))
        for x, c in zip(excel.records, csv.records):
            for f in fields(TerrainRecord):
                if f.name in self.SOURCE_ONLY:
                    continue
                with self.subTest(terreno=x.terreno, campo=f.name):
                    left, right = getattr(x, f.name), getattr(c, f.name)
                    if isinstance(left, float) or isinstance(right, float):
                        self.assertAlmostEqual(left, right)
                    else:
                        self.assertEqual(left, right)
            self.assertEqual(x.clave_dedupe, c.clave_dedupe)
        self.assertEqual({k: [f.codigo for f in v] for k, v in validate_all(excel.records).items()},
                         {k: [f.codigo for f in v] for k, v in validate_all(csv.records).items()})


class FakePostgres:
    """Just enough of postgres.session for Staging's pending_import table."""

    def __init__(self):
        self.rows = {}

    @contextmanager
    def session(self):
        yield self

    def execute(self, sql, params=()):
        if sql.startswith("INSERT INTO pending_import"):
            token, payload, created = params
            json.loads(payload)  # must be real JSON
            self.rows[token] = {"payload": payload, "created": created}
        elif "WHERE token" in sql:
            return _Cursor(self.rows.pop(params[0], None))
        return _Cursor(None)


class _Cursor:
    def __init__(self, row):
        self.row = row

    def fetchone(self):
        return self.row


class CloudStaging(unittest.TestCase):
    def setUp(self):
        self.fake = FakePostgres()
        patches = [patch("server.postgres.enabled", return_value=True),
                   patch("server.postgres.session", self.fake.session)]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    def csv_result(self):
        text = (CSV_TEXT + "4,,Jalisco,Tala,,,,,,,,,\n"
                "5,Predio Raro,Jalisco,Tala,\"Camino\nlargo\",x,,,,,,\"19,4326\",-99\n")
        return read_csv_bytes(text.encode())

    def test_round_trip_keeps_everything(self):
        result = self.csv_result()
        self.assertTrue(result.rechazadas and any(r.notes for r in result.records))
        incidencias = validate_all(result.records)

        staging = Staging()
        pending = staging.put("t.csv", "t", result, incidencias)
        self.assertEqual(len(self.fake.rows), 1)
        recovered = staging.take(pending.token)

        self.assertEqual(recovered.resultado, result)
        self.assertEqual(recovered.incidencias, incidencias)
        raro = next(r for r in recovered.resultado.records if r.terreno == "Predio Raro")
        self.assertEqual(raro.fila, 6)
        self.assertEqual({n.valor for n in raro.notes}, {"x", "19,4326"})
        self.assertTrue(all(n.motivo for n in raro.notes))
        self.assertIsNone(staging.take(pending.token))  # consumed once

    def test_extras_survive(self):
        result = read_csv_bytes(b"Terreno,Vendedor\nPredio A,Fulano\n")
        staging = Staging()
        recovered = staging.take(staging.put("t.csv", "t", result, validate_all(result.records)).token)
        self.assertEqual(dict(recovered.resultado.records[0].extra), {"Vendedor": "Fulano"})

    def test_payloads_from_before_motivo_still_decode(self):
        result = read_csv_bytes(b"Terreno,Asking Price\nPredio A,SD\n")
        staging = Staging()
        pending = staging.put("t.csv", "t", result, validate_all(result.records))
        stored = json.loads(self.fake.rows[pending.token]["payload"])
        for note in stored["resultado"]["records"][0]["notes"]:
            note.pop("motivo")
        self.fake.rows[pending.token]["payload"] = json.dumps(stored)
        recovered = staging.take(pending.token)
        self.assertEqual(recovered.resultado.records[0].notes[0].valor, "SD")


class ExcelRegression(NoCloud):
    def test_real_fixture_preview_is_unchanged(self):
        preview = self.preview(body=FIXTURE.read_bytes(), filename="Base Terrenos 09.26.xlsx")
        self.assertEqual((preview["conteo"], preview["ubicados"], preview["filas_con_datos"]), (79, 39, 79))
        self.assertEqual(preview["formato"], "xlsx")
        self.assertEqual(preview["resumen"]["VALOR_NO_NUMERICO"], 3)
        mensaje = next(i["mensaje"] for h in preview["hallazgos"] for i in h["incidencias"]
                       if i["codigo"] == "VALOR_NO_NUMERICO")
        self.assertEqual(mensaje, "«SD» no es un número; el campo asking_m2 quedó vacío.")


if __name__ == "__main__":
    unittest.main()
